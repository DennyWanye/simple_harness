"""(1) 看代表模板的 layout 名(判断 name-match 还是回退 index)
(2) 实测 A-2 用该模板生成一份 deck，重开报每页填了啥。结果写 UTF-8 文件。"""
import sys
from pathlib import Path
sys.path.insert(0, r"G:\projects\deskpet\backend")
from pptx import Presentation  # noqa: E402
from deskpet.tools.ppt_tools import ppt_create  # noqa: E402

OUT = Path(r"G:\projects\deskpet\plans\2026-06-09-ppt-beauty-mainline\render_probe_report.txt")
D = Path(r"G:\projects\deskpet\resources\PPT_Template")
samples = ["高级感 (01)---.pptx", "高级感 (28)---.pptx", "简约高级ppt-01（素材） (29).pptx"]
lines = []

for name in samples:
    f = D / name
    lines.append("=" * 80)
    lines.append(f"TEMPLATE: {name}")
    if not f.is_file():
        lines.append("  NOT FOUND"); continue
    prs = Presentation(str(f))
    lines.append(f"  layouts ({sum(len(m.slide_layouts) for m in prs.slide_masters)}):")
    for m in prs.slide_masters:
        for li, lay in enumerate(m.slide_layouts):
            phs = [str(ph.placeholder_format.type).split(' ')[0] for ph in lay.placeholders]
            lines.append(f"    [{li}] {lay.name!r}: {phs}")

    # 实测渲染
    outline = [
        {"layout": "title", "title": "人工智能简介", "subtitle": "DeskPet 自动生成测试"},
        {"layout": "bullet", "title": "三大要点", "bullets": ["机器学习", "深度神经网络", "大语言模型"]},
        {"layout": "section", "title": "应用场景", "subtitle": "落地领域概览"},
        {"layout": "two_column", "title": "对比", "left_title": "优势", "left": ["效率高", "可扩展"],
         "right_title": "挑战", "right": ["算力成本", "数据需求"]},
    ]
    outp = str(Path(r"G:\projects\deskpet\.tmp") / f"tpltest-{abs(hash(name))%10000}.pptx")
    res = ppt_create(outline, template=str(f), output_path=outp, title="测试")
    lines.append(f"  RENDER ok={res.get('ok')} path={res.get('path')} theme={res.get('theme')} slides={res.get('slide_count')}")
    if res.get("ok") and res.get("path") and Path(res["path"]).is_file():
        rp = Presentation(res["path"])
        for si, s in enumerate(rp.slides):
            texts = []
            for sh in s.shapes:
                if sh.has_text_frame and sh.text_frame.text.strip():
                    texts.append(sh.text_frame.text.strip()[:30])
            lines.append(f"    slide[{si}] layout={s.slide_layout.name!r} texts={texts}")

OUT.write_text("\n".join(lines), encoding="utf-8")
print("written", OUT, "lines", len(lines))
