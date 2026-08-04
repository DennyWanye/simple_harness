from __future__ import annotations

import json
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

from deskpet.tools import ppt_outline_store
from deskpet.tools import ppt_tools
from deskpet.memory.session_db import SessionDB
from deskpet.workflows import WorkflowContext
from deskpet.workflows.definitions.research_core import (
    FetchPort,
    ResearchLLMPort,
    ResearchSearchPort,
)
from deskpet.workflows.definitions.v1 import (
    PPT_PRO_V1,
    PPT_PRO_V1_DEFINITION,
    ppt_pro_initial_state,
)
from deskpet.workflows.definitions.v1 import ppt_pro as ppt_pro_graph
from deskpet.workflows.definitions import ppt_pro_nodes
from deskpet.workflows.errors import WorkflowErrorCode, WorkflowNodeError
from deskpet.workflows.outbox import WorkflowOutbox
from deskpet.workflows.progress import WorkflowProgressReporter
from deskpet.workflows.runner import WorkflowRegistry, WorkflowRunner
from deskpet.workflows.service import WorkflowService
from deskpet.workflows.store import FencedAsyncSqliteSaver, WorkflowRunStore
from tests.test_workflow_native_engine import MemoryNativeStore


class FakeLLM:
    def __init__(self, responses: list[str]) -> None:
        self.responses = list(responses)
        self.calls: list[str] = []

    async def complete(self, prompt: str) -> str:
        self.calls.append(prompt)
        if not self.responses:
            raise AssertionError("unexpected LLM call")
        return self.responses.pop(0)


class FakeToolPort:
    def __init__(
        self,
        root: Path,
        *,
        render_ok: bool = True,
        crash_slide: str = "",
        stable_previews: bool = False,
        stable_preview_pages: set[int] | None = None,
    ) -> None:
        self.root = root
        self.render_ok = render_ok
        self.crash_slide = crash_slide
        self.crashed = False
        self.image_calls: list[str] = []
        self.render_calls = 0
        self.render_modes: list[str] = []
        self.last_slide_count = 0
        self.stable_previews = stable_previews
        self.stable_preview_pages = set(stable_preview_pages or set())

    async def probe_images(self, *, timeout_s: float) -> bool:
        assert timeout_s > 0
        return True

    async def generate_slide_image(self, **kwargs):
        slide_id = kwargs["slide_id"]
        self.image_calls.append(slide_id)
        if slide_id == self.crash_slide and not self.crashed:
            self.crashed = True
            raise WorkflowNodeError(
                code=WorkflowErrorCode.RETRYABLE_PROVIDER,
                message_ref="test:image_crash",
                node_id="image_map",
            )
        path = self.root / f"{slide_id}.png"
        path.write_bytes(slide_id.encode("utf-8"))
        return {"path": str(path)}

    async def render_ppt(self, **kwargs):
        self.render_calls += 1
        self.render_modes.append(str(kwargs.get("render_mode") or ""))
        self.last_slide_count = len(kwargs.get("slides") or [])
        if not self.render_ok:
            return {"ok": False, "error": "render failed"}
        path = self.root / f"deck-{kwargs['render_revision']}.pptx"
        deck = Presentation()
        while len(deck.slides) < self.last_slide_count:
            deck.slides.add_slide(deck.slide_layouts[6])
        deck.save(path)
        return {
            "ok": True,
            "path": str(path),
            "slide_count": self.last_slide_count,
            "artifacts": [],
        }

    async def render_preview(self, **kwargs):
        previews = []
        for page in range(1, self.last_slide_count + 1):
            path = self.root / f"preview-{self.render_calls}-{page}.png"
            revision = (
                0
                if self.stable_previews or page in self.stable_preview_pages
                else self.render_calls
            )
            path.write_bytes(f"preview-{revision}-{page}".encode())
            sha256 = ppt_pro_nodes.file_or_value_hash(str(path), {})
            previews.append(
                {
                    "kind": "image",
                    "path": str(path),
                    "page": page,
                    "sha256": sha256,
                    "artifact_ref": f"preview:{sha256}",
                    "render_hash": str(kwargs.get("render_hash") or ""),
                }
            )
        return previews


class FakeEvaluator:
    def __init__(self, reviews: list[list[dict]] | None = None) -> None:
        self.reviews = list(reviews or [[]])
        self.calls = 0

    async def evaluate_ppt(self, **kwargs):
        self.calls += 1
        issues = self.reviews.pop(0) if self.reviews else []
        issue_pages = {
            int(issue.get("page") or 0)
            for issue in issues
            if isinstance(issue, dict)
        }
        reviews = [
            {
                "page": int(preview.get("page") or index),
                "ok": int(preview.get("page") or index) not in issue_pages,
            }
            for index, preview in enumerate(kwargs.get("previews") or [], start=1)
        ]
        return {
            "issues": issues,
            "reviews": reviews,
            "reviewed_pages": len(reviews),
            "score": 1.0 if not issues else 0.5,
        }


def _preview_ref(path: Path, page: int, *, render_hash: str = "render") -> dict:
    sha256 = ppt_pro_nodes.file_or_value_hash(str(path), {})
    return {
        "page": page,
        "path": str(path),
        "sha256": sha256,
        "artifact_ref": f"preview:{sha256}",
        "render_hash": render_hash,
    }


def test_rendered_deck_quality_gate_rejects_corruption_and_wrong_page_count(
    tmp_path: Path,
) -> None:
    valid_path = tmp_path / "valid.pptx"
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(valid_path)

    assert ppt_pro_nodes.inspect_rendered_deck(
        str(valid_path), expected_slide_count=2
    )["passed"]
    mismatch = ppt_pro_nodes.inspect_rendered_deck(
        str(valid_path), expected_slide_count=3
    )
    assert not mismatch["passed"]
    assert mismatch["hard_failures"][0]["code"] == "ppt_slide_count_mismatch"

    corrupt_path = tmp_path / "corrupt.pptx"
    corrupt_path.write_bytes(b"not-a-pptx")
    corrupt = ppt_pro_nodes.inspect_rendered_deck(
        str(corrupt_path), expected_slide_count=2
    )
    assert not corrupt["passed"]
    assert corrupt["hard_failures"][0]["code"] == "ppt_package_invalid"

    truncated_path = tmp_path / "truncated.pptx"
    with zipfile.ZipFile(truncated_path, "w") as archive:
        archive.writestr(
            "ppt/presentation.xml",
            (
                '<p:presentation xmlns:p="http://schemas.openxmlformats.org/'
                'presentationml/2006/main"><p:sldIdLst>'
                '<p:sldId id="256"/><p:sldId id="257"/>'
                "</p:sldIdLst></p:presentation>"
            ),
        )
    truncated = ppt_pro_nodes.inspect_rendered_deck(
        str(truncated_path), expected_slide_count=2
    )
    assert not truncated["passed"]
    assert truncated["hard_failures"][0]["code"] == "ppt_package_parts_missing"

    missing_rels_path = tmp_path / "missing-slide-rels.pptx"
    with zipfile.ZipFile(valid_path) as source, zipfile.ZipFile(
        missing_rels_path, "w"
    ) as target:
        for item in source.infolist():
            if item.filename == "ppt/slides/_rels/slide1.xml.rels":
                continue
            target.writestr(item, source.read(item.filename))
    missing_rels = ppt_pro_nodes.inspect_rendered_deck(
        str(missing_rels_path), expected_slide_count=2
    )
    assert not missing_rels["passed"]
    assert "ppt_slide_relationships_missing" in {
        item["code"] for item in missing_rels["hard_failures"]
    }


@pytest.mark.asyncio
async def test_visual_gate_rejects_missing_preview_before_calling_evaluator(
    tmp_path: Path,
) -> None:
    preview = tmp_path / "page-1.png"
    preview.write_bytes(b"preview")

    class MustNotRun:
        async def evaluate_ppt(self, **_kwargs):
            raise AssertionError("evaluator must not run on incomplete previews")

    result = await ppt_pro_nodes.evaluate_visuals(
        [{"page": 1, "path": str(preview)}],
        [{"slide_id": "one"}, {"slide_id": "two"}],
        MustNotRun(),
    )

    assert result["hard_failures"][0]["code"] == "ppt_preview_coverage_incomplete"


@pytest.mark.asyncio
async def test_visual_gate_rejects_incomplete_review_coverage(tmp_path: Path) -> None:
    previews = []
    for page in (1, 2):
        path = tmp_path / f"page-{page}.png"
        path.write_bytes(b"preview")
        previews.append(_preview_ref(path, page))

    class PartialEvaluator:
        async def evaluate_ppt(self, **_kwargs):
            return {
                "issues": [],
                "reviews": [{"page": 1, "ok": True}],
                "reviewed_pages": 1,
                "score": 1.0,
            }

    result = await ppt_pro_nodes.evaluate_visuals(
        previews,
        [{"slide_id": "one"}, {"slide_id": "two"}],
        PartialEvaluator(),
    )

    assert [failure["code"] for failure in result["hard_failures"]] == [
        "ppt_visual_review_coverage_incomplete"
    ]


@pytest.mark.asyncio
async def test_visual_gate_rejects_out_of_order_previews(tmp_path: Path) -> None:
    previews = []
    for page in (2, 1):
        path = tmp_path / f"page-{page}.png"
        path.write_bytes(b"preview")
        previews.append(_preview_ref(path, page))

    class MustNotRun:
        async def evaluate_ppt(self, **_kwargs):
            raise AssertionError("evaluator must not run on out-of-order previews")

    result = await ppt_pro_nodes.evaluate_visuals(
        previews,
        [{"slide_id": "one"}, {"slide_id": "two"}],
        MustNotRun(),
    )

    assert result["hard_failures"][0]["code"] == "ppt_preview_coverage_incomplete"


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "payload, expected_code",
    [
        (
            {
                "issues": [],
                "reviews": [{"page": 1.5, "ok": True}],
                "reviewed_pages": 1,
            },
            "ppt_visual_reviews_invalid",
        ),
        (
            {
                "issues": "MALFORMED",
                "reviews": [{"page": 1, "ok": True}],
                "reviewed_pages": 1,
            },
            "ppt_visual_review_issues_invalid",
        ),
        (
            {
                "issues": [],
                "reviews": [{"page": 1, "ok": True}],
                "reviewed_pages": 1,
                "hard_failures": ["MALFORMED"],
            },
            "ppt_visual_review_hard_failures_invalid",
        ),
        (
            {
                "issues": [],
                "reviews": [{"page": 1, "ok": True}],
                "reviewed_pages": 1,
                "hard_failures": "",
            },
            "ppt_visual_review_hard_failures_invalid",
        ),
        (
            {
                "issues": [],
                "reviews": [{"page": 1, "ok": True}],
                "reviewed_pages": 1,
                "hard_failures": {},
            },
            "ppt_visual_review_hard_failures_invalid",
        ),
    ],
)
async def test_visual_gate_rejects_malformed_evaluator_payloads(
    tmp_path: Path, payload: dict, expected_code: str
) -> None:
    preview = tmp_path / "page-1.png"
    preview.write_bytes(b"preview")

    class MalformedEvaluator:
        async def evaluate_ppt(self, **_kwargs):
            return payload

    result = await ppt_pro_nodes.evaluate_visuals(
        [_preview_ref(preview, 1)],
        [{"slide_id": "one"}],
        MalformedEvaluator(),
    )

    assert expected_code in {
        item["code"] for item in result["hard_failures"]
    }


@pytest.mark.asyncio
async def test_publish_revalidates_missing_render_artifact(tmp_path: Path) -> None:
    deck_path = tmp_path / "gone.pptx"
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(deck_path)
    output_hash = ppt_pro_nodes.file_or_value_hash(str(deck_path), {})
    preview = tmp_path / "page-1.png"
    preview.write_bytes(b"preview")
    deck_path.unlink()

    patch = await ppt_pro_graph.publish_handler(
        {
            "values": {
                "render_ref": {
                    "ok": True,
                    "path": str(deck_path),
                    "output_hash": output_hash,
                },
                "slide_records": [{"slide_id": "one"}],
                "preview_refs": [{"page": 1, "path": str(preview)}],
                "visual_review": {
                    "issues": [],
                    "reviews": [{"page": 1, "ok": True}],
                    "reviewed_pages": 1,
                    "hard_failures": [],
                },
            }
        },
        WorkflowContext(),
    )

    assert patch.values["values"]["terminal_status"] == "error"
    assert "publish_ref" not in patch.values["values"]
    assert "artifact_refs" not in patch.values


@pytest.mark.asyncio
async def test_full_page_preview_rejects_non_integer_source_page(tmp_path: Path) -> None:
    page_image = tmp_path / "page.png"
    page_image.write_bytes(b"page")

    refs = await ppt_pro_nodes.render_preview(
        {
            "render_mode": ppt_pro_nodes.FULL_PAGE_RENDER_MODE,
            "output_hash": "render",
            "page_images": [{"page": True, "path": str(page_image)}],
        },
        None,
    )

    assert refs == []


async def _publish_values(tmp_path: Path) -> tuple[dict, Path]:
    deck_path = tmp_path / "deck.pptx"
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(deck_path)
    output_hash = ppt_pro_nodes.file_or_value_hash(str(deck_path), {})
    preview_path = tmp_path / "page-1.png"
    preview_path.write_bytes(b"original-preview")
    preview_ref = _preview_ref(preview_path, 1, render_hash=output_hash)
    review = {
        "issues": [],
        "reviews": [{"page": 1, "ok": True}],
        "reviewed_pages": 1,
        "score": 1.0,
        "hard_failures": [],
    }
    review["review_hash"] = ppt_pro_nodes.content_hash(review)
    return (
        {
            "render_ref": {
                "ok": True,
                "path": str(deck_path),
                "output_hash": output_hash,
            },
            "slide_records": [{"slide_id": "one"}],
            "preview_refs": [preview_ref],
            "preview_hash": ppt_pro_nodes.content_hash([preview_ref]),
            "visual_review": review,
        },
        preview_path,
    )


@pytest.mark.asyncio
async def test_publish_rejects_replaced_preview_and_forged_review_hash(
    tmp_path: Path,
) -> None:
    values, preview_path = await _publish_values(tmp_path)
    valid = await ppt_pro_graph.publish_handler(
        {"values": values}, WorkflowContext()
    )
    assert valid.values["values"]["terminal_status"] == "success"

    preview_path.write_bytes(b"replaced-after-review")
    replaced = await ppt_pro_graph.publish_handler(
        {"values": values}, WorkflowContext()
    )
    assert replaced.values["values"]["terminal_status"] == "error"
    assert "artifact_refs" not in replaced.values

    restored, _ = await _publish_values(tmp_path)
    restored["visual_review"]["review_hash"] = "FORGED"
    forged = await ppt_pro_graph.publish_handler(
        {"values": restored}, WorkflowContext()
    )
    assert forged.values["values"]["terminal_status"] == "error"
    assert "artifact_refs" not in forged.values


@pytest.mark.asyncio
async def test_publish_rejects_consistent_but_unresolved_visual_issue(
    tmp_path: Path,
) -> None:
    values, _ = await _publish_values(tmp_path)
    review = {
        "issues": [{"page": 1, "ok": False}],
        "reviews": [{"page": 1, "ok": False}],
        "reviewed_pages": 1,
        "score": 0.0,
        "hard_failures": [],
    }
    review["review_hash"] = ppt_pro_nodes.content_hash(review)
    values["visual_review"] = review

    patch = await ppt_pro_graph.publish_handler(
        {"values": values}, WorkflowContext()
    )

    assert patch.values["values"]["terminal_status"] == "error"
    assert "ppt_visual_issues_remaining" in patch.values["values"]["terminal_error"]
    assert "artifact_refs" not in patch.values


def _research_config() -> dict[str, object]:
    return {
        "max_sub_questions": 2,
        "max_urls_per_query": 1,
        "max_total_passages": 2,
        "min_passage_chars": 20,
        "max_rounds": 1,
        "query_expansion": False,
        "site_directed": False,
        "source_packs": False,
        "direct_sources": False,
        "rerank_mode": "off",
    }


def _outline(title: str = "Deck") -> str:
    return json.dumps(
        [
            {
                "layout": "image_full",
                "title": title,
                "bullets": ["one", "two", "three"],
                "image_prompt": "cinematic cover, no text",
            },
            {
                "layout": "image_full",
                "title": "Evidence",
                "bullets": ["alpha", "beta", "gamma"],
                "image_prompt": "cinematic evidence, no text",
            },
        ]
    )


def _llm(*extra_outlines: str) -> FakeLLM:
    return FakeLLM(
        [
            json.dumps(["what evidence supports the topic?"]),
            "# Research\n\n## TL;DR\n\nSupported finding [^1].",
            _outline(),
            *extra_outlines,
        ]
    )


class _ProgressRecorder:
    def __init__(self) -> None:
        self.events: list[tuple[str, str]] = []

    async def report(self, identity, transition):
        self.events.append((identity.node_id, transition))


def _context(
    llm: FakeLLM,
    tool: FakeToolPort,
    evaluator: FakeEvaluator,
    *,
    notifier=None,
    progress=None,
) -> WorkflowContext:
    async def search(query: str, *, max_results: int):
        return [{"url": "https://example.com/source", "title": query}]

    async def fetch(url: str):
        return {
            "ok": True,
            "url": url,
            "title": "Source",
            "text": "durable research evidence " * 40,
            "fetched_at": 123.0,
        }

    ports = {
            "llm": ResearchLLMPort(complete=llm.complete),
            "search": ResearchSearchPort(search_call=search),
            "fetch": FetchPort(extractor=fetch),
            "tool": tool,
            "evaluator": evaluator,
    }
    if notifier is not None:
        ports["notifier"] = notifier
    if progress is not None:
        ports["progress"] = progress
    return WorkflowContext(ports=ports)


def _state(tmp_path: Path, run_id: str, **overrides):
    values = {
        "topic": "Durable decks",
        "run_id": run_id,
        "session_id": "session",
        "pages": 3,
        "full_page_images": True,
        "research_config": _research_config(),
        "blob_root": tmp_path / "blobs",
        **overrides,
    }
    return ppt_pro_initial_state(**values)


@pytest.mark.asyncio
@pytest.mark.parametrize("pages", [1, 2])
async def test_ppt_graph_normalize_preserves_short_deck_page_count(
    tmp_path: Path, pages: int
) -> None:
    state = _state(tmp_path, f"short-{pages}", pages=pages)
    context = _context(_llm(), FakeToolPort(tmp_path), FakeEvaluator())
    patch = await ppt_pro_graph.normalize_handler(state, context)

    assert patch.values["values"]["pages"] == pages


@pytest.mark.asyncio
async def test_prepare_slides_typed_layout_failure_routes_to_safe_terminal(
    tmp_path: Path, monkeypatch
) -> None:
    def fail(*args, **kwargs):
        raise ppt_tools.FullPageLayoutError("text_overflow")

    monkeypatch.setattr(ppt_pro_graph.nodes, "stable_slide_records", fail)
    state = _state(tmp_path, "layout-failure")
    values = dict(state["values"])
    values.update(
        {
            "outline_slides": [{"title": "Too much"}],
            "image_mode": True,
            "full_page_images": True,
            "image_probe": {"requested": True, "reachable": True},
        }
    )
    state["values"] = values

    patch = await ppt_pro_graph.prepare_slides_handler(state, WorkflowContext())
    failed_state = dict(state)
    failed_state["values"] = patch.values["values"]

    assert patch.values["values"]["terminal_status"] == "error"
    assert patch.values["values"]["terminal_error"]["code"] == "text_overflow"
    assert await ppt_pro_graph.prepare_slides_route(failed_state, WorkflowContext()) == "terminal"
    assert "traceback" not in str(patch.values).lower()


@pytest.mark.asyncio
async def test_prepare_slides_removes_flattened_layouts_for_editable_run(
    tmp_path: Path,
) -> None:
    state = _state(
        tmp_path,
        "editable-layout-normalization",
        pages=4,
        full_page_images=False,
    )
    state["values"].update(
        {
            "editable_required": True,
            "outline_slides": [
                {
                    "layout": "image_full",
                    "title": "Evidence",
                    "bullets": ["Native evidence text"],
                    "image_prompt": "abstract evidence background",
                },
                {
                    "layout": "quote",
                    "title": "Risks",
                    "bullets": ["Risk one", "Risk two"],
                },
                {
                    "layout": "two_column",
                    "title": "Architecture",
                    "bullets": ["Run", "Kernel", "Driver", "Tool"],
                },
                {
                    "layout": "section",
                    "title": "Timeline",
                    "bullets": ["Find", "Fix", "Verify"],
                },
            ],
            "image_mode": True,
            "image_probe": {"requested": True, "reachable": True},
        }
    )

    patch = await ppt_pro_graph.prepare_slides_handler(state, WorkflowContext())
    records = patch.values["values"]["slide_records"]

    assert [record["slide"]["layout"] for record in records] == [
        "image",
        "bullet",
        "two_column",
        "toc",
    ]
    assert records[0]["slide"]["bullets"] == ["Native evidence text"]
    assert records[1]["slide"]["bullets"] == ["Risk one", "Risk two"]
    assert records[2]["slide"]["left"] == ["Run", "Kernel"]
    assert records[2]["slide"]["right"] == ["Driver", "Tool"]
    assert records[3]["slide"]["bullets"] == ["Find", "Fix", "Verify"]


async def _invoke_to_waiting(
    tmp_path, *, llm=None, tool=None, evaluator=None, notifier=None, progress=None, **state_overrides
):
    run_id = "ppt-run"
    saver = MemoryNativeStore()
    executable = PPT_PRO_V1.bind(checkpointer=saver)
    llm = llm or _llm()
    tool = tool or FakeToolPort(tmp_path)
    evaluator = evaluator or FakeEvaluator()
    context = _context(llm, tool, evaluator, notifier=notifier, progress=progress)
    waiting = await executable.ainvoke(
        _state(tmp_path, run_id, **state_overrides),
        context,
        thread_id=run_id,
        run_id=run_id,
    )
    return executable, saver, context, waiting, llm, tool, evaluator


def _interrupt_id(waiting: dict) -> str:
    interrupt = waiting["interrupt"]
    return str(interrupt["interrupt_id"])


@pytest.mark.asyncio
async def test_compiled_ppt_progress_persists_to_original_session(
    tmp_path, monkeypatch
) -> None:
    monkeypatch.setenv("DESKPET_STATE_DB_PATH", str(tmp_path / "state.db"))
    session_id = "ppt-progress-session"
    session_db = SessionDB(tmp_path / "session.db")
    await session_db.initialize()
    await session_db.ensure_session(session_id)
    workflow_db = tmp_path / "workflow.db"
    store = WorkflowRunStore(workflow_db)
    registry = WorkflowRegistry()
    registry.register(PPT_PRO_V1)
    runner = WorkflowRunner(
        store,
        FencedAsyncSqliteSaver(workflow_db),
        registry,
        owner="progress-ppt",
    )
    run_id = await runner.start(
        session_id=session_id,
        request_id="ppt-progress-request",
        turn_id="ppt-progress-turn",
        workflow_name="ppt_pro",
        workflow_version="v1",
        capability_snapshot={"ports": ["llm", "search", "fetch", "tool"]},
    )
    await store.bind_session_refs(run_id, (("delivery", session_id, 0),))

    async def persist(event, delivery) -> None:
        await session_db.append_message_if_epoch(
            delivery["target_id"],
            "assistant",
            str(event["payload"]["text"]),
            expected_epoch=0,
            workflow_event_id=str(event["event_id"]),
        )

    outbox = WorkflowOutbox(store)
    service = WorkflowService(
        store,
        runner=runner,
        outbox=outbox,
        delivery_handlers={"session_message": persist},
    )
    progress = WorkflowProgressReporter(
        service, (("session_message", session_id),)
    )
    result = await runner.run(
        run_id,
        _state(tmp_path, run_id, session_id=session_id),
        _context(_llm(), FakeToolPort(tmp_path), FakeEvaluator(), progress=progress),
    )

    assert result.status.value == "waiting"
    messages = await session_db.get_messages(session_id)
    event_ids = [message["workflow_event_id"] for message in messages]
    assert len(event_ids) == len(set(event_ids))
    events = await outbox.events_after(run_id, 0, limit=100)
    ordinals = {
        event["payload"]["ordinal"]
        for event in events["events"]
        if event["event_type"] == "workflow.progress"
    }
    assert ordinals.issuperset({1, 2, 3, 4, 5, 6})


def _intents(output: dict) -> list[dict]:
    return output["values"]["delivery_intents"]


@pytest.mark.asyncio
async def test_outline_ready_notifies_original_session_before_waiting(
    tmp_path, monkeypatch
) -> None:
    monkeypatch.setenv("DESKPET_STATE_DB_PATH", str(tmp_path / "state.db"))
    notifications = []

    async def notify(payload):
        notifications.append(payload)

    progress = _ProgressRecorder()
    await _invoke_to_waiting(tmp_path, notifier=notify, progress=progress)

    assert len(notifications) == 1
    assert notifications[0]["run_id"] == "ppt-run"
    assert notifications[0]["session_id"] == "session"
    assert notifications[0]["outline_id"].startswith("workflow:ppt-run:")
    assert "Durable decks" in notifications[0]["topic"]
    assert notifications[0]["outline_md"]
    started = {node for node, transition in progress.events if transition == "started"}
    assert {"normalize", "research_plan", "research_search", "research_synth", "outline", "wait_outline_decision"}.issubset(started)


@pytest.mark.asyncio
async def test_legacy_checkpoint_without_full_page_flag_resumes_in_images_mode(
    tmp_path, monkeypatch
) -> None:
    monkeypatch.setenv("DESKPET_STATE_DB_PATH", str(tmp_path / "state.db"))
    run_id = "legacy-ppt-run"
    saver = MemoryNativeStore()
    executable = PPT_PRO_V1.bind(checkpointer=saver)
    llm = _llm()
    tool = FakeToolPort(tmp_path)
    context = _context(llm, tool, FakeEvaluator())
    state = _state(tmp_path, run_id)
    values = dict(state["values"])
    values.pop("full_page_images", None)
    values.pop("editable_required", None)
    values.pop("slide_requirements", None)
    state["values"] = values

    waiting = await executable.ainvoke(
        state,
        context,
        thread_id=run_id,
        run_id=run_id,
    )
    restarted = PPT_PRO_V1.bind(checkpointer=saver)
    completed = await restarted.resume(
        {_interrupt_id(waiting): {"action": "accept"}},
        context,
        thread_id=run_id,
        run_id=run_id,
    )

    artifact = next(item for item in _intents(completed) if item["kind"] == "artifact_card")
    assert artifact["payload"]["render_mode"] == "images"
    assert tool.render_modes == ["images"]


@pytest.mark.asyncio
async def test_outline_restart_continue_does_not_repeat_research_or_outline(
    tmp_path, monkeypatch
) -> None:
    monkeypatch.setenv("DESKPET_STATE_DB_PATH", str(tmp_path / "state.db"))
    executable, saver, context, waiting, llm, tool, _ = await _invoke_to_waiting(tmp_path)

    assert "delivery_intents" not in waiting.get("values", {})
    assert len(llm.calls) == 3
    pending_id = _interrupt_id(waiting)

    restarted = PPT_PRO_V1.bind(checkpointer=saver)
    completed = await restarted.resume(
        {pending_id: {"action": "accept"}},
        context,
        thread_id="ppt-run",
        run_id="ppt-run",
    )

    assert len(llm.calls) == 3
    assert len(tool.image_calls) == 2
    assert {item["kind"] for item in _intents(completed)} == {
        "receipt",
        "artifact_card",
        "open_artifact",
        "final_assistant",
    }
    assert completed["artifact_refs"][0].startswith("pptx:")
    assert completed["receipt_refs"] == ["ppt-run:receipt"]


@pytest.mark.asyncio
async def test_outline_generic_approval_response_maps_to_accept(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("DESKPET_STATE_DB_PATH", str(tmp_path / "state.db"))
    executable, saver, context, waiting, llm, tool, _ = await _invoke_to_waiting(tmp_path)

    restarted = PPT_PRO_V1.bind(checkpointer=saver)
    completed = await restarted.resume(
        {_interrupt_id(waiting): {"approved": True}},
        context,
        thread_id="ppt-run",
        run_id="ppt-run",
    )

    assert len(llm.calls) == 3
    assert len(tool.image_calls) == 2
    assert completed["artifact_refs"]


@pytest.mark.asyncio
async def test_slide_pending_recovery_reuses_committed_image_checkpoint(
    tmp_path, monkeypatch
) -> None:
    monkeypatch.setenv("DESKPET_STATE_DB_PATH", str(tmp_path / "state.db"))
    tool = FakeToolPort(tmp_path)
    executable, saver, context, waiting, _, tool, _ = await _invoke_to_waiting(
        tmp_path, tool=tool
    )
    pending_id = _interrupt_id(waiting)
    # Stable ids are visible in the accepted outline and survive preparation.
    tool.crash_slide = "slide-002-" + "placeholder"

    original_generate = tool.generate_slide_image

    async def crash_second(**kwargs):
        if len(tool.image_calls) == 1 and not tool.crashed:
            tool.crashed = True
            tool.image_calls.append(kwargs["slide_id"])
            raise WorkflowNodeError(
                code=WorkflowErrorCode.RETRYABLE_PROVIDER,
                message_ref="test:image_crash",
                node_id="image_map",
            )
        return await original_generate(**kwargs)

    tool.generate_slide_image = crash_second  # type: ignore[method-assign]
    with pytest.raises(WorkflowNodeError):
        await executable.resume(
            {pending_id: {"action": "accept"}},
            context,
            thread_id="ppt-run",
            run_id="ppt-run",
        )

    first_slide_id = tool.image_calls[0]
    restarted = PPT_PRO_V1.bind(checkpointer=saver)
    completed = await restarted.ainvoke(
        None,
        context,
        thread_id="ppt-run",
        run_id="ppt-run",
    )

    assert completed["values"]["delivery_intents"]
    assert tool.image_calls.count(first_slide_id) == 1
    assert len(set(tool.image_calls)) == 2


@pytest.mark.asyncio
async def test_outline_revision_budget_is_persisted_and_bounded(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("DESKPET_STATE_DB_PATH", str(tmp_path / "state.db"))
    llm = _llm(_outline("Revision 1"), _outline("Revision 2"))
    executable, _, context, waiting, llm, _, _ = await _invoke_to_waiting(
        tmp_path, llm=llm, max_outline_revisions=2
    )

    for feedback in ("first", "second"):
        waiting = await executable.resume(
            {_interrupt_id(waiting): {"action": "modify", "feedback": feedback}},
            context,
            thread_id="ppt-run",
            run_id="ppt-run",
        )
        assert "interrupt" in waiting

    completed = await executable.resume(
        {_interrupt_id(waiting): {"action": "modify", "feedback": "third"}},
        context,
        thread_id="ppt-run",
        run_id="ppt-run",
    )

    receipt = next(item for item in _intents(completed) if item["kind"] == "receipt")
    assert receipt["payload"]["outcome"] == "error"
    assert "budget exhausted" in receipt["payload"]["error"]
    assert len(llm.calls) == 5


@pytest.mark.asyncio
async def test_outline_modify_emits_new_card_and_does_not_start_generation(
    tmp_path, monkeypatch
) -> None:
    monkeypatch.setenv("DESKPET_STATE_DB_PATH", str(tmp_path / "state.db"))
    notifications = []

    async def notify(payload):
        notifications.append(payload)

    llm = _llm(_outline("Revised deck"))
    executable, _, context, waiting, _, tool, _ = await _invoke_to_waiting(
        tmp_path, llm=llm, notifier=notify
    )
    original_outline_id = notifications[-1]["outline_id"]

    revised_waiting = await executable.resume(
        {
            _interrupt_id(waiting): {
                "action": "modify",
                "feedback": "把第二页改成时间线",
            }
        },
        context,
        thread_id="ppt-run",
        run_id="ppt-run",
    )

    assert "interrupt" in revised_waiting
    assert len(notifications) == 2
    assert notifications[-1]["outline_id"] != original_outline_id
    assert notifications[-1]["outline_md"]
    assert tool.image_calls == []
    assert tool.render_calls == 0


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("action", "render_ok", "outcome", "has_artifact"),
    [
        ("accept", True, "success", True),
        ("accept", False, "error", False),
        ("cancel", True, "cancelled", False),
    ],
)
async def test_terminal_intents_cover_success_error_cancel_without_research_delivery(
    tmp_path, monkeypatch, action, render_ok, outcome, has_artifact
) -> None:
    monkeypatch.setenv("DESKPET_STATE_DB_PATH", str(tmp_path / "state.db"))
    tool = FakeToolPort(tmp_path, render_ok=render_ok)
    executable, _, context, waiting, _, _, _ = await _invoke_to_waiting(
        tmp_path, tool=tool
    )
    assert "delivery_intents" not in waiting.get("values", {})

    completed = await executable.resume(
        {_interrupt_id(waiting): {"action": action}},
        context,
        thread_id="ppt-run",
        run_id="ppt-run",
    )
    intents = _intents(completed)
    receipt = next(item for item in intents if item["kind"] == "receipt")

    assert receipt["payload"]["outcome"] == outcome
    assert any(item["kind"] == "artifact_card" for item in intents) is has_artifact
    assert not any(item["kind"] == "report" for item in intents)
    assert tool.render_calls == (1 if action == "accept" else 0)


@pytest.mark.asyncio
async def test_full_page_provider_unavailable_fails_instead_of_using_template(
    tmp_path, monkeypatch
) -> None:
    monkeypatch.setenv("DESKPET_STATE_DB_PATH", str(tmp_path / "state.db"))
    tool = FakeToolPort(tmp_path)

    async def unavailable(*, timeout_s: float) -> bool:
        assert timeout_s > 0
        return False

    tool.probe_images = unavailable  # type: ignore[method-assign]
    executable, _, context, waiting, _, tool, _ = await _invoke_to_waiting(
        tmp_path, tool=tool
    )
    completed = await executable.resume(
        {_interrupt_id(waiting): {"action": "accept"}},
        context,
        thread_id="ppt-run",
        run_id="ppt-run",
    )

    assert tool.image_calls == []
    assert tool.render_calls == 0
    assert not any(item["kind"] == "artifact_card" for item in _intents(completed))
    receipt = next(item for item in _intents(completed) if item["kind"] == "receipt")
    assert receipt["payload"]["outcome"] == "error"
    final = next(item for item in _intents(completed) if item["kind"] == "final_assistant")
    assert "普通模板" in final["payload"]["text"]
    assert "稍后重试" in final["payload"]["text"]


@pytest.mark.asyncio
async def test_full_page_visual_revision_budget_fails_closed(
    tmp_path, monkeypatch
) -> None:
    monkeypatch.setenv("DESKPET_STATE_DB_PATH", str(tmp_path / "state.db"))
    issue = {"page": 1, "action": "shrink_text"}
    evaluator = FakeEvaluator([[issue], [issue], [issue]])
    executable, _, context, waiting, _, tool, evaluator = await _invoke_to_waiting(
        tmp_path, evaluator=evaluator, max_visual_revisions=2
    )

    completed = await executable.resume(
        {_interrupt_id(waiting): {"action": "accept"}},
        context,
        thread_id="ppt-run",
        run_id="ppt-run",
    )

    assert next(item for item in _intents(completed) if item["kind"] == "receipt")[
        "payload"
    ]["outcome"] == "error"
    assert not any(item["kind"] == "artifact_card" for item in _intents(completed))
    final = next(item for item in _intents(completed) if item["kind"] == "final_assistant")
    assert "PPT page quality issues remain" in final["payload"]["text"]
    assert completed["loop_counters"]["visual_revisions"] == 2
    assert evaluator.calls == 3
    assert tool.render_calls == 3
    first_slide_id, second_slide_id = tool.image_calls[:2]
    assert tool.image_calls.count(first_slide_id) == 3
    assert tool.image_calls.count(second_slide_id) == 1


@pytest.mark.asyncio
async def test_visual_revision_that_does_not_change_pixels_fails_immediately(
    tmp_path, monkeypatch
) -> None:
    monkeypatch.setenv("DESKPET_STATE_DB_PATH", str(tmp_path / "state.db"))
    issue = {"page": 1, "action": "shrink_text"}
    evaluator = FakeEvaluator([[issue], [issue]])
    tool = FakeToolPort(tmp_path, stable_previews=True)
    executable, _, context, waiting, _, tool, evaluator = await _invoke_to_waiting(
        tmp_path,
        tool=tool,
        evaluator=evaluator,
        max_visual_revisions=2,
    )

    completed = await executable.resume(
        {_interrupt_id(waiting): {"action": "accept"}},
        context,
        thread_id="ppt-run",
        run_id="ppt-run",
    )

    final = next(item for item in _intents(completed) if item["kind"] == "final_assistant")
    assert "did not change reviewed page(s): 1" in final["payload"]["text"]
    assert completed["values"]["visual_review"]["hard_failures"] == [
        {"code": "ppt_visual_revision_no_effect", "pages": [1]}
    ]
    assert completed["loop_counters"]["visual_revisions"] == 1
    assert evaluator.calls == 1
    assert tool.render_calls == 2


@pytest.mark.asyncio
async def test_visual_revision_fails_when_one_reviewed_page_does_not_change(
    tmp_path, monkeypatch
) -> None:
    monkeypatch.setenv("DESKPET_STATE_DB_PATH", str(tmp_path / "state.db"))
    issues = [
        {"page": 1, "action": "shrink_text"},
        {"page": 2, "action": "change_variant"},
    ]
    evaluator = FakeEvaluator([issues, issues])
    tool = FakeToolPort(tmp_path, stable_preview_pages={2})
    executable, _, context, waiting, _, tool, evaluator = await _invoke_to_waiting(
        tmp_path,
        tool=tool,
        evaluator=evaluator,
        max_visual_revisions=2,
    )

    completed = await executable.resume(
        {_interrupt_id(waiting): {"action": "accept"}},
        context,
        thread_id="ppt-run",
        run_id="ppt-run",
    )

    final = next(item for item in _intents(completed) if item["kind"] == "final_assistant")
    assert "did not change reviewed page(s): 2" in final["payload"]["text"]
    assert completed["values"]["visual_review"]["hard_failures"] == [
        {"code": "ppt_visual_revision_no_effect", "pages": [2]}
    ]
    assert completed["loop_counters"]["visual_revisions"] == 1
    assert evaluator.calls == 1
    assert tool.render_calls == 2


def test_manifest_topology_uses_durable_barrier_and_stable_slide_map() -> None:
    definition = PPT_PRO_V1_DEFINITION
    node_ids = [node.node_id for node in definition.nodes]
    barrier = next(node for node in definition.nodes if node.node_id == "wait_outline_decision")

    assert node_ids[:2] == ["normalize", "research_plan"]
    assert node_ids[-7:] == [
        "image_map",
        "render",
        "preview",
        "visual_evaluate",
        "visual_revise",
        "publish",
        "terminal",
    ]
    assert barrier.interrupt_capable and barrier.barrier and barrier.exclusive_superstep
    assert definition.policy_manifest["research_terminal_delivery"] is False
    assert definition.policy_manifest["slide_map_key"] == "stable_slide_id"
    assert definition.loop_budgets == {
        "gap_iterations": 1,
        "outline_revisions": 2,
        "visual_revisions": 2,
        "image_iterations": 120,
    }


@pytest.mark.asyncio
async def test_ppt_tool_uses_durable_workflow_starter() -> None:
    seen = []

    async def starter(payload):
        seen.append(payload)
        return {"run_id": "run-accepted"}

    old = ppt_tools._ppt_workflow_starter
    try:
        ppt_tools.set_ppt_pro_services(workflow_starter=starter)
        result = await ppt_tools._submit_ppt_pro(
            {"topic": "Accepted deck", "_session_id": "session"}, "call-1"
        )
    finally:
        ppt_tools.set_ppt_pro_services(workflow_starter=old)

    assert result["status"] == "accepted"
    assert result["run_id"] == "run-accepted"
    assert result["call_id"] == "call-1"
    assert seen[0]["workflow_name"] == "ppt_pro"


@pytest.mark.asyncio
async def test_ppt_tool_has_no_legacy_kill_switch() -> None:
    started = []

    async def starter(payload):
        started.append(payload)
        return {"run_id": "durable-run"}

    old = ppt_tools._ppt_workflow_starter
    try:
        ppt_tools.set_ppt_pro_services(workflow_starter=starter)
        result = await ppt_tools._submit_ppt_pro(
            {"topic": "Durable deck", "_session_id": "session"}, "call-2"
        )
    finally:
        ppt_tools.set_ppt_pro_services(workflow_starter=old)

    assert result["status"] == "accepted"
    assert result["run_id"] == "durable-run"
    assert len(started) == 1


def test_outline_history_restart_expiry_skips_graph_projection(tmp_path) -> None:
    database = tmp_path / "state.db"

    def connect():
        import sqlite3

        return sqlite3.connect(database)

    assert ppt_outline_store.save_outline(
        "legacy-outline", "session", "Legacy", [], 0, conn_factory=connect
    )
    graph_id = ppt_outline_store.workflow_outline_id("run-1", 0)
    assert ppt_outline_store.project_workflow_outline(
        graph_id, "session", "Graph", [], 0, conn_factory=connect
    )

    assert ppt_outline_store.expire_dangling_proposed(conn_factory=connect) == 1
    assert ppt_outline_store.get_outline(
        "legacy-outline", conn_factory=connect
    )["status"] == "expired"
    assert ppt_outline_store.get_outline(graph_id, conn_factory=connect)["status"] == "proposed"
