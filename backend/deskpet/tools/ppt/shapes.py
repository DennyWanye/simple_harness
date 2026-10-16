# SPDX-FileCopyrightText: 2026 DennyWanye
# SPDX-License-Identifier: BUSL-1.1

"""python-pptx drawing primitives: fills, shapes, text blocks, cards and badges."""

from __future__ import annotations

import re
from collections.abc import Iterable, Sequence

from .common import MSO_ANCHOR, MSO_SHAPE, PP_ALIGN, Emu, Inches, Pt, RGBColor
from .outline import Theme
from .research import _SLIDE_HEIGHT


def _rgb(triple: tuple[int, int, int]):
    return RGBColor(*triple)


def _fill_slide_bg(
    slide,
    theme: Theme,
    color: tuple[int, int, int] | None = None,
) -> None:
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(color or theme.background_rgb)


def _estimate_capacity(width: int, height: int, font_size: int) -> int:
    width_in = max(width / 914400, 0.5)
    height_in = max(height / 914400, 0.25)
    chars_per_line = max(int(width_in * 12.5 * 16 / max(font_size, 8)), 8)
    lines = max(int(height_in * 72 / max(font_size * 1.3, 1)), 1)
    return max(chars_per_line * lines, 12)


def _shorten(text: str, max_chars: int) -> str:
    text = str(text or "").strip()
    if len(text) <= max_chars:
        return text
    return text[: max(max_chars - 3, 1)].rstrip() + "..."


def _fit_font_size(text: str, width: int, height: int, desired: int, minimum: int = 10) -> int:
    size = desired
    while size > minimum and len(str(text)) > _estimate_capacity(width, height, size):
        size -= 1
    return size


def _add_shape(
    slide,
    shape_type,
    *,
    left: int,
    top: int,
    width: int,
    height: int,
    fill: tuple[int, int, int],
    line: tuple[int, int, int] | None = None,
):
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = _rgb(line)
    return shape


def _set_fill_alpha(shape, alpha_pct: int) -> None:
    """给已 solid 填充的形状加透明度(0=全透,100=不透)。python-pptx 无原生
    alpha API,直接往 srgbClr 注入 <a:alpha>。半透明深色遮罩让底图透出来
    (引用页/卡片可读又不死黑)。从不抛异常。"""
    try:
        # 按需导入：重型包
        from pptx.oxml.ns import qn

        srgb = shape.fill.fore_color._xFill.find(qn("a:srgbClr"))
        if srgb is None:
            return
        for old in srgb.findall(qn("a:alpha")):
            srgb.remove(old)
        a = srgb.makeelement(qn("a:alpha"), {"val": str(int(max(0, min(100, alpha_pct)) * 1000))})
        srgb.append(a)
    except Exception:  # noqa: BLE001
        pass


def _add_text(
    slide,
    text: str,
    *,
    left: int,
    top: int,
    width: int,
    height: int,
    font_size: int,
    bold: bool = False,
    color: tuple[int, int, int] = (0, 0, 0),
    font_name: str = "Microsoft YaHei UI",
    align: str = "left",
    anchor: str = "top",
    margin: float = 0.05,
):
    font_size = _fit_font_size(text, width, height, font_size)
    text = _shorten(text, _estimate_capacity(width, height, font_size))
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = {
        "top": MSO_ANCHOR.TOP,
        "middle": MSO_ANCHOR.MIDDLE,
        "bottom": MSO_ANCHOR.BOTTOM,
    }.get(anchor, MSO_ANCHOR.TOP)
    p = tf.paragraphs[0]
    p.alignment = {
        "left": PP_ALIGN.LEFT,
        "center": PP_ALIGN.CENTER,
        "right": PP_ALIGN.RIGHT,
    }.get(align, PP_ALIGN.LEFT)
    run = p.add_run()
    run.text = text
    f = run.font
    f.name = font_name
    f.size = Pt(font_size)
    f.bold = bold
    f.color.rgb = _rgb(color)
    return tb


def _add_bullet_text(
    slide,
    bullets: Iterable[str],
    *,
    left: int,
    top: int,
    width: int,
    height: int,
    font_size: int,
    color: tuple[int, int, int],
    font_name: str,
    accent_color: tuple[int, int, int] | None = None,
):
    """Render bullets with a coloured dot prefix.

    python-pptx doesn't expose first-class bullet formatting reliably
    across themes, so we draw the bullet character ourselves (●)
    coloured with the accent colour.
    """
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.1)
    tf.margin_right = Inches(0.1)
    tf.margin_top = Inches(0.05)
    tf.margin_bottom = Inches(0.05)
    bullets = list(bullets)
    if not bullets:
        # Empty — leave an invisible placeholder so layout stays stable.
        tf.paragraphs[0].text = ""
        return tb
    total_chars = sum(len(str(line)) for line in bullets)
    if total_chars > 1200:
        font_size = min(font_size, 9)
    elif total_chars > 850:
        font_size = min(font_size, 10)
    elif total_chars > 650:
        font_size = min(font_size, 11)
    elif total_chars > 480:
        font_size = min(font_size, 12)
    paragraph_spacing = 8 if len(bullets) <= 4 else 4
    for i, line in enumerate(bullets):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = PP_ALIGN.LEFT
        para.space_after = Pt(paragraph_spacing)
        dot = para.add_run()
        dot.text = "● "
        dot.font.size = Pt(font_size)
        dot.font.name = font_name
        dot.font.color.rgb = _rgb(accent_color or color)
        body = para.add_run()
        # Contract-relevant editable prose must never be silently rewritten
        # with an ellipsis. Dense layouts choose a smaller font above instead.
        body.text = str(line)
        body.font.size = Pt(font_size)
        body.font.name = font_name
        body.font.color.rgb = _rgb(color)
    return tb


def _add_title_block(slide, title: str, subtitle: str, theme: Theme) -> None:
    title_size = 36
    if len(title) > 24:
        title_size = 30
    if len(title) > 34:
        title_size = 26
    _add_text(
        slide,
        title,
        left=Inches(0.62),
        top=Inches(0.30),
        width=Inches(7.0),
        height=Inches(0.90),
        font_size=title_size,
        bold=True,
        color=theme.primary_rgb,
        font_name=theme.font_heading,
    )
    if subtitle:
        _add_text(
            slide,
            subtitle,
            left=Inches(0.64),
            top=Inches(0.95),
            width=Inches(7.7),
            height=Inches(0.36),
            font_size=12,
            color=theme.muted_text_rgb,
            font_name=theme.font_body,
        )


def _add_corner_motif(slide, theme: Theme, *, dark: bool = False) -> None:
    base = theme.secondary_rgb if dark else theme.surface_rgb
    _add_shape(
        slide,
        MSO_SHAPE.ROUNDED_RECTANGLE,
        left=Inches(7.78),
        top=Inches(0.34),
        width=Inches(1.45),
        height=Inches(0.82),
        fill=base,
    )
    _add_shape(
        slide,
        MSO_SHAPE.OVAL,
        left=Inches(8.92),
        top=Inches(0.70),
        width=Inches(0.34),
        height=Inches(0.34),
        fill=theme.accent_rgb,
    )


def _add_icon_badge(
    slide,
    label: str,
    *,
    left: int,
    top: int,
    size: int,
    theme: Theme,
    fill: tuple[int, int, int] | None = None,
    text_color: tuple[int, int, int] | None = None,
):
    _add_shape(
        slide,
        MSO_SHAPE.OVAL,
        left=left,
        top=top,
        width=size,
        height=size,
        fill=fill or theme.secondary_rgb,
    )
    return _add_text(
        slide,
        label,
        left=left,
        top=top + Inches(0.01),
        width=size,
        height=size,
        font_size=10,
        bold=True,
        color=text_color or (255, 255, 255),
        font_name="Calibri",
        align="center",
        anchor="middle",
        margin=0.0,
    )


_NUMBER_RE = re.compile(r"(?<![\w.^])(\d+(?:\.\d+)?\s?(?:%|x|X|倍|万|亿|年|小时|分钟|人|页)?)")


def _extract_callouts(
    items: Sequence[str], limit: int = 2
) -> tuple[list[tuple[str, str]], list[str]]:
    callouts: list[tuple[str, str]] = []
    remaining: list[str] = []
    for item in items:
        text = str(item).strip()
        match = _NUMBER_RE.search(text)
        if match and len(callouts) < limit:
            value = match.group(1).replace(" ", "")
            label = (text[: match.start()] + text[match.end() :]).strip(" :-：，,。")
            count_prefix = text[: match.start()]
            count_suffix = text[match.end() :]
            is_parenthetical_count = (
                re.search(r"[(（]\s*$", count_prefix) is not None
                and re.match(
                    r"\s*(?:项|条|个|页|类|items?|entries?|rows?)?\s*[)）]",
                    count_suffix,
                    flags=re.IGNORECASE,
                )
                is not None
            )
            # A count embedded in prose, such as "缺陷（1项）", describes the
            # sentence. Turning it into a compact metric card can truncate the
            # required editable copy that follows it.
            # Two callouts use the compact 1.46-inch card. PowerPoint's real
            # text bounds show that more than 30 mixed Chinese/Latin
            # characters need a third line and collide with the metric. Keep
            # longer evidence in the wide editable body layout instead.
            if is_parenthetical_count or len(label) > 30:
                remaining.append(text)
                continue
            callouts.append((value, label or text))
        else:
            remaining.append(text)
    return callouts, remaining


def _add_callout_card(
    slide,
    value: str,
    label: str,
    *,
    left: int,
    top: int,
    width: int,
    height: int,
    theme: Theme,
):
    _add_shape(
        slide,
        MSO_SHAPE.ROUNDED_RECTANGLE,
        left=left,
        top=top,
        width=width,
        height=height,
        fill=theme.primary_rgb,
    )
    _add_shape(
        slide,
        MSO_SHAPE.OVAL,
        left=left + width - Inches(0.58),
        top=top + Inches(0.18),
        width=Inches(0.32),
        height=Inches(0.32),
        fill=theme.accent_rgb,
    )
    _add_text(
        slide,
        value,
        left=left + Inches(0.28),
        top=top + Inches(0.12),
        width=width - Inches(0.55),
        height=Inches(0.62),
        font_size=34,
        bold=True,
        color=theme.dark_text_rgb,
        font_name="Calibri",
        anchor="middle",
        margin=0.0,
    )
    _add_text(
        slide,
        label,
        left=left + Inches(0.32),
        top=top + Inches(0.78),
        width=width - Inches(0.64),
        height=height - Inches(0.86),
        font_size=12,
        color=theme.dark_muted_text_rgb,
        font_name=theme.font_body,
        anchor="middle",
        margin=0.04,
    )


def _add_item_card(
    slide,
    title: str,
    body: str,
    *,
    left: int,
    top: int,
    width: int,
    height: int,
    theme: Theme,
    badge: str,
):
    _add_shape(
        slide,
        MSO_SHAPE.ROUNDED_RECTANGLE,
        left=left,
        top=top,
        width=width,
        height=height,
        fill=theme.surface_rgb,
        line=(226, 232, 240),
    )
    if height <= Inches(1.25):
        _add_icon_badge(
            slide,
            badge,
            left=left + Inches(0.18),
            top=top + Inches(0.18),
            size=Inches(0.34),
            theme=theme,
        )
        _add_text(
            slide,
            body,
            left=left + Inches(0.68),
            top=top + Inches(0.12),
            width=width - Inches(0.90),
            height=height - Inches(0.24),
            font_size=10 if len(body) > 44 else 12,
            color=theme.text_rgb,
            font_name=theme.font_body,
            anchor="middle",
        )
        return
    _add_icon_badge(
        slide,
        badge,
        left=left + Inches(0.22),
        top=top + Inches(0.22),
        size=Inches(0.42),
        theme=theme,
    )
    _add_text(
        slide,
        title,
        left=left + Inches(0.78),
        top=top + Inches(0.20),
        width=width - Inches(1.0),
        height=Inches(0.34),
        font_size=16,
        bold=True,
        color=theme.primary_rgb,
        font_name=theme.font_heading,
    )
    _add_text(
        slide,
        body,
        left=left + Inches(0.28),
        top=top + Inches(0.76),
        width=width - Inches(0.56),
        height=height - Inches(0.88),
        font_size=14,
        color=theme.text_rgb,
        font_name=theme.font_body,
    )


def _add_accent_bar(slide, theme: Theme, *, top: int = 0) -> None:
    """Decorative left-edge accent bar for non-title slides."""
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Emu(top),
        Inches(0.18),
        Emu(_SLIDE_HEIGHT - top),
    )
    bar.line.fill.background()
    bar.fill.solid()
    bar.fill.fore_color.rgb = _rgb(theme.accent_rgb)
