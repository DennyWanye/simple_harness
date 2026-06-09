"""探查 python-pptx 默认模板的布局/占位符，给 A-2 模板填充 spec 用。"""
import sys
sys.path.insert(0, r"G:\projects\deskpet\backend")
from pptx import Presentation  # noqa: E402

prs = Presentation()  # 默认内置模板
print(f"#slides={len(prs.slides)}  #layouts={len(prs.slide_layouts)}")
for i, lay in enumerate(prs.slide_layouts):
    phs = [(ph.placeholder_format.idx, str(ph.placeholder_format.type), ph.name) for ph in lay.placeholders]
    print(f"layout[{i}] {lay.name!r}: {phs}")
