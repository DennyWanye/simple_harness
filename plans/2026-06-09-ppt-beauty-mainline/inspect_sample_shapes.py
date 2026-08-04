"""精查生成样例每页所有形状: 是否占位符/类型/文本/位置,搞清装饰vs填充。"""
import sys
from pathlib import Path
sys.path.insert(0, r"G:\projects\deskpet\backend")
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

OUT = Path(r"G:\projects\deskpet\plans\2026-06-09-ppt-beauty-mainline\sample_shapes_report.txt")
f = r"G:\projects\deskpet\.tmp\ppt-samples\样例A-高级感01.pptx"
prs = Presentation(f)
lines = [f"FILE: {f}  slides={len(prs.slides)}"]
for si, s in enumerate(prs.slides):
    lines.append(f"\n=== slide[{si}] layout={s.slide_layout.name!r} ===")
    for sh in s.shapes:
        is_ph = getattr(sh, "is_placeholder", False)
        pht = ""
        if is_ph:
            try:
                pht = str(sh.placeholder_format.type).split(" ")[0]
            except Exception:
                pht = "?"
        txt = ""
        if sh.has_text_frame:
            txt = sh.text_frame.text.replace("\n", " / ")[:50]
        try:
            pos = f"L{Emu(sh.left).inches:.1f} T{Emu(sh.top).inches:.1f} W{Emu(sh.width).inches:.1f}"
        except Exception:
            pos = "pos?"
        lines.append(f"  ph={is_ph}({pht:<11}) {pos:<22} txt={txt!r}")
OUT.write_text("\n".join(lines), encoding="utf-8")
print("written", OUT)
