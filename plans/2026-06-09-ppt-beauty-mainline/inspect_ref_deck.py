"""探查参考 deck 的母版/布局结构，判断能否当模板复用。"""
import sys
from pathlib import Path
sys.path.insert(0, r"G:\projects\deskpet\backend")
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

d = r"G:\projects\deskpet\.tmp\ai-education-deck.pptx"
prs = Presentation(d)
print(f"slide_size = {Emu(prs.slide_width).inches:.2f} x {Emu(prs.slide_height).inches:.2f} in")
print(f"#slides = {len(prs.slides)}")
print(f"#masters = {len(prs.slide_masters)}")
for mi, master in enumerate(prs.slide_masters):
    print(f"\n=== master[{mi}] '{master.name}'  #layouts={len(master.slide_layouts)} ===")
    for li, lay in enumerate(master.slide_layouts):
        phs = [(ph.placeholder_format.idx, str(ph.placeholder_format.type), ph.name) for ph in lay.placeholders]
        print(f"  layout[{li}] name={lay.name!r}  placeholders={phs}")

# 各 slide 用的 layout
print("\n=== 现有 slides 用的 layout ===")
for i, s in enumerate(prs.slides):
    print(f"  slide[{i}] layout={s.slide_layout.name!r}")
