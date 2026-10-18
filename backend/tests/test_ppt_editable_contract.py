from __future__ import annotations

import zipfile

import pytest
from deskpet.tools import ppt_tools
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

from tests.ppt_patching import patch_ppt


@pytest.mark.parametrize("length", [79, 80, 81, 160])
def test_template_text_replacement_preserves_complete_long_copy(length: int) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1.5))
    paragraph = shape.text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = "template"
    run.font.size = Pt(18)
    expected = "长" * length

    ppt_tools._set_text_keep_style(shape, [expected])

    assert shape.text == expected
    assert len(shape.text) == length


def test_editable_visual_revision_changes_real_slide_output(tmp_path) -> None:
    image_path = tmp_path / "background.png"
    Image.new("RGB", (1600, 900), (12, 32, 72)).save(image_path)
    plain_path = tmp_path / "plain.pptx"
    revised_path = tmp_path / "revised.pptx"
    common = {
        "layout": "bullet",
        "title": "Harness revision",
        "bullets": ["First long editable point", "Second long editable point"],
        "image_path": str(image_path),
    }

    assert ppt_tools.ppt_create(
        [common],
        output_path=str(plain_path),
        skip_image_gen=True,
        force_fromscratch=True,
    )["ok"]
    assert ppt_tools.ppt_create(
        [{**common, "image_variant": "top", "font_scale": 0.81}],
        output_path=str(revised_path),
        skip_image_gen=True,
        force_fromscratch=True,
    )["ok"]

    with zipfile.ZipFile(plain_path) as archive:
        plain_xml = archive.read("ppt/slides/slide1.xml")
    with zipfile.ZipFile(revised_path) as archive:
        revised_xml = archive.read("ppt/slides/slide1.xml")
    assert revised_xml != plain_xml

    revised = Presentation(revised_path)
    visible_text = "\n".join(
        shape.text for shape in revised.slides[0].shapes if getattr(shape, "has_text_frame", False)
    )
    assert "First long editable point" in visible_text
    assert "Second long editable point" in visible_text


def test_dense_chart_notes_get_room_and_visual_revision_never_enlarges_text(
    tmp_path,
) -> None:
    plain_path = tmp_path / "dense-chart-plain.pptx"
    revised_path = tmp_path / "dense-chart-revised.pptx"
    bullets = [
        "根据优先级矩阵对六项缺陷进行统计：P0/P1 共 1 项、P1 共 2 项、P2 共 3 项，确保修复路线合理性",
        "P0/P1 缺陷数量为 1：Voice 绕过 ProductTurnPreparer，兼具 P0 紧急性与 P1 重要性",
        "P1 缺陷数量为 2：核心模块过度膨胀与核心层依赖方向不干净，计划通过模块拆分与依赖倒置解决",
        "P2 缺陷数量为 3：投影一致性窗口、跨 Run 启发式指代与兼容代码混居，纳入长期技术债管理",
    ]
    common = {
        "layout": "chart",
        "title": "架构缺陷优先级分布（P0/P1、P1、P2）",
        "bullets": bullets,
        "chart": {
            "type": "bar",
            "categories": ["P0/P1", "P1", "P2"],
            "series": [{"name": "缺陷数量", "data": [1, 2, 3]}],
        },
    }

    assert ppt_tools.ppt_create(
        [common],
        output_path=str(plain_path),
        skip_image_gen=True,
        force_fromscratch=True,
    )["ok"]
    assert ppt_tools.ppt_create(
        [{**common, "font_scale": 0.9}],
        output_path=str(revised_path),
        skip_image_gen=True,
        force_fromscratch=True,
    )["ok"]

    def _notes_shape_and_sizes(path):
        presentation = Presentation(path)
        notes_shape = next(
            shape
            for shape in presentation.slides[0].shapes
            if getattr(shape, "has_text_frame", False) and bullets[0] in shape.text
        )
        sizes = [
            run.font.size.pt
            for paragraph in notes_shape.text_frame.paragraphs
            for run in paragraph.runs
            if run.font.size is not None
        ]
        return presentation, notes_shape, sizes

    plain_presentation, plain_notes, plain_sizes = _notes_shape_and_sizes(plain_path)
    _, revised_notes, revised_sizes = _notes_shape_and_sizes(revised_path)
    assert plain_notes.top.inches < 3
    assert plain_notes.height.inches >= 2.2
    assert plain_notes.top + plain_notes.height <= plain_presentation.slide_height
    assert revised_notes.height == plain_notes.height
    assert plain_sizes and revised_sizes
    assert max(revised_sizes) <= max(plain_sizes)
    assert min(revised_sizes) <= min(plain_sizes)


def test_editable_table_visual_revision_scales_cell_text(tmp_path) -> None:
    plain_path = tmp_path / "table-plain.pptx"
    revised_path = tmp_path / "table-revised.pptx"
    rows = [
        ["Priority", "Defect", "Impact", "Owner", "Status"],
        *[
            [
                f"P{index % 3}",
                f"Architecture defect {index}",
                "Dense editable explanation that may wrap",
                "Harness",
                "Fixed",
            ]
            for index in range(1, 7)
        ],
    ]
    common = {
        "layout": "table",
        "title": "Architecture defects",
        "table": {"header": True, "rows": rows},
    }

    assert ppt_tools.ppt_create(
        [common],
        output_path=str(plain_path),
        skip_image_gen=True,
        force_fromscratch=True,
    )["ok"]
    assert ppt_tools.ppt_create(
        [{**common, "font_scale": 0.81}],
        output_path=str(revised_path),
        skip_image_gen=True,
        force_fromscratch=True,
    )["ok"]

    plain_table = next(
        shape.table
        for shape in Presentation(plain_path).slides[0].shapes
        if getattr(shape, "has_table", False)
    )
    revised_table = next(
        shape.table
        for shape in Presentation(revised_path).slides[0].shapes
        if getattr(shape, "has_table", False)
    )

    def _cell_font_points(table) -> list[float]:
        return [
            run.font.size.pt
            for row in table.rows
            for cell in row.cells
            for paragraph in cell.text_frame.paragraphs
            for run in paragraph.runs
            if run.font.size is not None
        ]

    plain_sizes = _cell_font_points(plain_table)
    revised_sizes = _cell_font_points(revised_table)
    assert plain_sizes
    assert revised_sizes
    assert max(plain_sizes) == 13
    assert min(plain_sizes) == 12
    assert max(revised_sizes) == 10
    assert min(revised_sizes) == 9
    assert [cell.text for row in revised_table.rows for cell in row.cells] == [
        cell.text for row in plain_table.rows for cell in row.cells
    ]


@pytest.mark.parametrize("layout", ["chart", "toc"])
def test_change_variant_changes_chart_and_toc_slide_output(tmp_path, monkeypatch, layout) -> None:
    patch_ppt(monkeypatch, "_default_template", lambda: None)
    common = {
        "layout": layout,
        "title": f"{layout.title()} revision",
        "bullets": ["Alpha", "Beta", "Gamma", "Delta"],
    }
    if layout == "chart":
        common["chart"] = {
            "type": "bar",
            "categories": ["A", "B", "C", "D"],
            "series": [{"name": "Score", "data": [85, 90, 75, 80]}],
        }

    plain_path = tmp_path / f"{layout}-plain.pptx"
    revised_path = tmp_path / f"{layout}-revised.pptx"
    assert ppt_tools.ppt_create(
        [common],
        output_path=str(plain_path),
        skip_image_gen=True,
        force_fromscratch=True,
    )["ok"]
    assert ppt_tools.ppt_create(
        [{**common, "image_variant": "top"}],
        output_path=str(revised_path),
        skip_image_gen=True,
        force_fromscratch=True,
    )["ok"]

    with zipfile.ZipFile(plain_path) as archive:
        plain_xml = archive.read("ppt/slides/slide1.xml")
    with zipfile.ZipFile(revised_path) as archive:
        revised_xml = archive.read("ppt/slides/slide1.xml")
    assert revised_xml != plain_xml

    visible_text = "\n".join(
        shape.text
        for shape in Presentation(revised_path).slides[0].shapes
        if getattr(shape, "has_text_frame", False)
    )
    for bullet in common["bullets"]:
        assert bullet in visible_text
