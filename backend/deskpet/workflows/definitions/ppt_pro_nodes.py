# SPDX-FileCopyrightText: 2026 DennyWanye
# SPDX-License-Identifier: BUSL-1.1

"""Durable stage adapters shared by the PPT Pro v1 graph."""

from __future__ import annotations

import copy
import hashlib
import inspect
import json
import posixpath
import re
import zipfile
from dataclasses import asdict
from itertools import combinations
from pathlib import Path
from typing import Any, Mapping, Sequence
from urllib.parse import unquote
from xml.etree import ElementTree

from ...tools import ppt_tools
from ..contracts import JsonValue, canonical_json, validate_json_value


FULL_PAGE_RENDER_MODE = "full_page_images"
_FULL_PAGE_REVISION_CODES = frozenset(
    {
        "blank",
        "layout_misfit",
        "text_extra",
        "text_incorrect",
        "text_missing",
        "text_occluded",
        "text_overflow",
        "text_unreadable",
    }
)
_REVISION_INSTRUCTIONS = {
    "blank": "Fill the page with the requested composition and exact copy.",
    "layout_misfit": "Recompose the page for a clearer presentation hierarchy.",
    "text_extra": "Remove every word or character not present in the exact copy.",
    "text_incorrect": "Correct all rendered text to match the exact copy verbatim.",
    "text_missing": "Render every non-empty string from the exact copy.",
    "text_occluded": "Keep all text unobstructed with strong contrast.",
    "text_overflow": "Use larger safe margins and keep all text inside the canvas.",
    "text_unreadable": "Increase text size, spacing, and contrast for readability.",
}


def content_hash(value: JsonValue) -> str:
    validate_json_value(value)
    return hashlib.sha256(canonical_json(value).encode("utf-8")).hexdigest()


def file_or_value_hash(path: str | None, value: JsonValue) -> str:
    if path:
        candidate = Path(path)
        try:
            if candidate.is_file():
                digest = hashlib.sha256()
                with candidate.open("rb") as handle:
                    for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                        digest.update(chunk)
                return digest.hexdigest()
        except OSError:
            pass
    return content_hash(value)


def _normalized_ppt_text(value: object) -> str:
    return "".join(char.casefold() for char in str(value or "") if char.isalnum())


def _expected_editable_copy(record: Mapping[str, Any]) -> list[tuple[str, str]]:
    payload = record.get("slide")
    slide = payload if isinstance(payload, Mapping) else record
    layout = str(slide.get("layout") or "bullet")
    fields: list[tuple[str, object]] = [
        ("title", slide.get("title")),
        ("subtitle", slide.get("subtitle")),
    ]
    if layout in {"bullet", "image", "toc"}:
        fields.extend(("bullet", item) for item in slide.get("bullets") or ())
    elif layout == "two_column":
        fields.extend(
            [
                ("left_title", slide.get("left_title")),
                *((("left", item) for item in slide.get("left") or ())),
                ("right_title", slide.get("right_title")),
                *((("right", item) for item in slide.get("right") or ())),
            ]
        )
    elif layout == "quote":
        fields.extend(
            [("quote", slide.get("quote")), ("cite", slide.get("cite"))]
        )
    elif layout == "table":
        table = slide.get("table")
        if isinstance(table, Mapping):
            fields.extend(
                ("table_cell", cell)
                for row in table.get("rows") or ()
                if isinstance(row, Sequence) and not isinstance(row, (str, bytes))
                for cell in row
            )
    elif layout == "chart":
        fields.extend(("bullet", item) for item in slide.get("bullets") or ())
        chart = slide.get("chart")
        if isinstance(chart, Mapping):
            fields.extend(
                ("chart_category", item) for item in chart.get("categories") or ()
            )
            fields.extend(
                ("chart_series", series.get("name"))
                for series in chart.get("series") or ()
                if isinstance(series, Mapping)
            )
    return [
        (field, normalized)
        for field, value in fields
        if len(normalized := _normalized_ppt_text(value)) >= 2
    ]


# Digits in priority/version labels such as ``P0``/``P1``/``V2``/``R2`` are
# not detachable callout metrics. A single digit in natural-language copy
# (for example ``1项``), including English copy with spaces removed by the
# normalizer, is still a valid editable metric.
_EDITABLE_CALLOUT_METRIC_RE = re.compile(r"(?<![0-9prv])\d+")
_MAX_EDITABLE_SPLIT_METRICS = 12


def _editable_copy_present(
    expected_text: str,
    actual_native_text: Sequence[str],
) -> bool:
    if any(expected_text in item for item in actual_native_text):
        return True
    # The renderer may promote a percentage such as ``6.3%`` into its own
    # editable callout shape. Accept that split only when the exact metric is
    # a standalone native shape and the remaining prose is intact elsewhere.
    metric_matches = tuple(_EDITABLE_CALLOUT_METRIC_RE.finditer(expected_text))
    standalone_metrics = tuple(
        item
        for item in actual_native_text
        if _EDITABLE_CALLOUT_METRIC_RE.fullmatch(item)
    )
    if (
        not metric_matches
        or not standalone_metrics
        or len(metric_matches) > _MAX_EDITABLE_SPLIT_METRICS
    ):
        return False

    def ordered_subsequence(metrics: Sequence[str]) -> bool:
        cursor = 0
        for metric in metrics:
            while (
                cursor < len(standalone_metrics)
                and standalone_metrics[cursor] != metric
            ):
                cursor += 1
            if cursor >= len(standalone_metrics):
                return False
            cursor += 1
        return True

    # One numeric token may remain inline while another occurrence of the
    # same token becomes a callout. Enumerate the occurrences removed from
    # the prose instead of replacing every equal value at once; this preserves
    # multiplicity and rejects swapped multi-metric callouts.
    max_split = min(len(metric_matches), len(standalone_metrics))
    for split_count in range(1, max_split + 1):
        for split_indexes in combinations(range(len(metric_matches)), split_count):
            split_set = set(split_indexes)
            split_values = tuple(
                metric_matches[index].group(0) for index in split_indexes
            )
            if not ordered_subsequence(split_values):
                continue
            parts: list[str] = []
            cursor = 0
            for index, match in enumerate(metric_matches):
                if index not in split_set:
                    continue
                parts.append(expected_text[cursor : match.start()])
                cursor = match.end()
            parts.append(expected_text[cursor:])
            remaining = "".join(parts)
            if len(remaining) >= 2 and any(
                remaining in item for item in actual_native_text
            ):
                return True
    return False


def inspect_rendered_deck(
    path: str | None,
    *,
    expected_slide_count: int,
    editable_required: bool = False,
    slide_requirements: Mapping[str, Sequence[str]] | None = None,
    expected_slides: Sequence[Mapping[str, Any]] | None = None,
) -> dict[str, JsonValue]:
    """Fail-closed structural check before a rendered deck may reach preview."""

    failures: list[dict[str, JsonValue]] = []
    actual_slide_count: int | None = None
    shape_inventory: list[dict[str, JsonValue]] = []
    try:
        candidate = Path(path).expanduser().resolve() if path else None
        candidate_exists = bool(candidate and candidate.is_file())
        candidate_size = candidate.stat().st_size if candidate_exists else 0
    except OSError:
        candidate = None
        candidate_exists = False
        candidate_size = 0
    if not candidate_exists:
        failures.append({"code": "ppt_file_missing"})
    elif candidate is None or candidate.suffix.casefold() != ".pptx" or candidate_size <= 0:
        failures.append({"code": "ppt_file_invalid"})
    else:
        try:
            with zipfile.ZipFile(candidate) as archive:
                corrupt_member = archive.testzip()
                if corrupt_member:
                    failures.append(
                        {
                            "code": "ppt_zip_corrupt",
                            "member": str(corrupt_member),
                        }
                    )
                presentation = ElementTree.fromstring(
                    archive.read("ppt/presentation.xml")
                )
                namespace = {
                    "p": "http://schemas.openxmlformats.org/presentationml/2006/main"
                }
                slide_ids = presentation.findall(".//p:sldIdLst/p:sldId", namespace)
                actual_slide_count = len(slide_ids)
                required_members = {
                    "[Content_Types].xml",
                    "_rels/.rels",
                    "ppt/presentation.xml",
                    "ppt/_rels/presentation.xml.rels",
                }
                missing_members = sorted(required_members.difference(archive.namelist()))
                if missing_members:
                    failures.append(
                        {
                            "code": "ppt_package_parts_missing",
                            "members": missing_members,
                        }
                    )
                member_names = set(archive.namelist())
                slide_parts = sorted(
                    name
                    for name in member_names
                    if name.startswith("ppt/slides/slide")
                    and name.endswith(".xml")
                    and "/_rels/" not in name
                )
                missing_slide_rels = [
                    (
                        f"{posixpath.dirname(name)}/_rels/"
                        f"{posixpath.basename(name)}.rels"
                    )
                    for name in slide_parts
                    if (
                        f"{posixpath.dirname(name)}/_rels/"
                        f"{posixpath.basename(name)}.rels"
                    )
                    not in member_names
                ]
                if missing_slide_rels:
                    failures.append(
                        {
                            "code": "ppt_slide_relationships_missing",
                            "members": missing_slide_rels,
                        }
                    )
                relationship_namespace = {
                    "r": (
                        "http://schemas.openxmlformats.org/package/2006/"
                        "relationships"
                    )
                }
                missing_targets: list[str] = []
                for rels_name in sorted(
                    name for name in member_names if name.endswith(".rels")
                ):
                    rels = ElementTree.fromstring(archive.read(rels_name))
                    if "/_rels/" in rels_name:
                        source_dir = rels_name.rsplit("/_rels/", 1)[0]
                    elif rels_name == "_rels/.rels":
                        source_dir = ""
                    else:
                        source_dir = posixpath.dirname(rels_name)
                    for relation in rels.findall("r:Relationship", relationship_namespace):
                        if str(relation.get("TargetMode") or "").casefold() == "external":
                            continue
                        target = unquote(str(relation.get("Target") or "")).split("#", 1)[0]
                        if not target:
                            missing_targets.append(f"{rels_name}:<empty>")
                            continue
                        normalized = posixpath.normpath(
                            posixpath.join(source_dir, target.lstrip("/"))
                        )
                        if normalized not in member_names:
                            missing_targets.append(f"{rels_name}:{normalized}")
                if missing_targets:
                    failures.append(
                        {
                            "code": "ppt_relationship_target_missing",
                            "targets": missing_targets,
                        }
                    )
        except (KeyError, OSError, ElementTree.ParseError, zipfile.BadZipFile) as exc:
            failures.append(
                {
                    "code": "ppt_package_invalid",
                    "detail": exc.__class__.__name__,
                }
            )
        if not failures:
            try:
                from pptx import Presentation

                loaded = Presentation(str(candidate))
                actual_slide_count = len(loaded.slides)
                requirements = slide_requirements or {}
                for raw_page in requirements:
                    page_text = str(raw_page)
                    try:
                        required_page = int(page_text)
                    except (TypeError, ValueError):
                        required_page = 0
                    if (
                        required_page < 1
                        or required_page > expected_slide_count
                        or page_text != str(required_page)
                    ):
                        failures.append(
                            {
                                "code": "ppt_required_page_out_of_range",
                                "page": page_text,
                            }
                        )
                if (
                    editable_required
                    and expected_slides is not None
                    and len(expected_slides) != expected_slide_count
                ):
                    failures.append(
                        {
                            "code": "ppt_expected_editable_records_mismatch",
                            "expected": expected_slide_count,
                            "actual": len(expected_slides),
                        }
                    )
                for page, slide in enumerate(loaded.slides, start=1):
                    pictures = 0
                    text_shapes = 0
                    visible_text_shapes = 0
                    visible_text_entries: list[tuple[int, int, str]] = []
                    native_object_text: list[str] = []
                    tables = 0
                    charts = 0
                    chart_types: list[str] = []
                    for shape in slide.shapes:
                        try:
                            if int(shape.shape_type) == 13:
                                pictures += 1
                        except (TypeError, ValueError):
                            pass
                        if bool(getattr(shape, "has_text_frame", False)):
                            shape_text = str(getattr(shape, "text", "") or "").strip()
                            if shape_text:
                                text_shapes += 1
                                normalized_text = _normalized_ppt_text(shape_text)
                                left = int(getattr(shape, "left", 0) or 0)
                                top = int(getattr(shape, "top", 0) or 0)
                                right = left + int(getattr(shape, "width", 0) or 0)
                                bottom = top + int(getattr(shape, "height", 0) or 0)
                                visible_width = max(
                                    0, min(right, loaded.slide_width) - max(left, 0)
                                )
                                visible_height = max(
                                    0, min(bottom, loaded.slide_height) - max(top, 0)
                                )
                                visible_area = visible_width * visible_height
                                slide_area = loaded.slide_width * loaded.slide_height
                                if (
                                    (
                                        len(normalized_text) >= 2
                                        or bool(
                                            _EDITABLE_CALLOUT_METRIC_RE.fullmatch(
                                                normalized_text
                                            )
                                        )
                                    )
                                    and visible_width >= loaded.slide_width * 0.01
                                    and visible_height >= loaded.slide_height * 0.01
                                    and visible_area >= slide_area * 0.0005
                                ):
                                    visible_text_shapes += 1
                                    visible_text_entries.append(
                                        (top, left, normalized_text)
                                    )
                        if bool(getattr(shape, "has_table", False)):
                            tables += 1
                            for row in shape.table.rows:
                                for cell in row.cells:
                                    normalized_cell = _normalized_ppt_text(cell.text)
                                    if normalized_cell:
                                        native_object_text.append(normalized_cell)
                        if bool(getattr(shape, "has_chart", False)):
                            charts += 1
                            chart = shape.chart
                            chart_types.append(str(chart.chart_type).casefold())
                            for series in chart.series:
                                normalized_name = _normalized_ppt_text(series.name)
                                if normalized_name:
                                    native_object_text.append(normalized_name)
                            for plot in chart.plots:
                                for category in plot.categories:
                                    normalized_category = _normalized_ppt_text(
                                        getattr(category, "label", category)
                                    )
                                    if normalized_category:
                                        native_object_text.append(
                                            normalized_category
                                        )
                    visible_text = [
                        item[2]
                        for item in sorted(
                            visible_text_entries,
                            key=lambda item: (item[0], item[1]),
                        )
                    ]
                    inventory: dict[str, JsonValue] = {
                        "page": page,
                        "pictures": pictures,
                        "text_shapes": text_shapes,
                        "visible_text_shapes": visible_text_shapes,
                        "tables": tables,
                        "charts": charts,
                        "chart_types": chart_types,
                    }
                    shape_inventory.append(inventory)
                    if editable_required and visible_text_shapes == 0:
                        failures.append(
                            {
                                "code": (
                                    "ppt_flattened_slide_forbidden"
                                    if pictures
                                    else "ppt_editable_content_missing"
                                ),
                                "page": page,
                            }
                        )
                    expected_record = (
                        expected_slides[page - 1]
                        if expected_slides and page <= len(expected_slides)
                        else {}
                    )
                    if editable_required and expected_record:
                        expected_payload = expected_record.get("slide")
                        expected_slide = (
                            expected_payload
                            if isinstance(expected_payload, Mapping)
                            else expected_record
                        )
                        if str(expected_slide.get("layout") or "") == "image_full":
                            failures.append(
                                {
                                    "code": "ppt_editable_layout_forbidden",
                                    "page": page,
                                    "layout": "image_full",
                                }
                            )
                        actual_native_text = [*visible_text, *native_object_text]
                        for field, expected_text in _expected_editable_copy(
                            expected_record
                        ):
                            if not _editable_copy_present(
                                expected_text,
                                actual_native_text,
                            ):
                                failures.append(
                                    {
                                        "code": (
                                            "ppt_expected_editable_title_missing"
                                            if field == "title"
                                            else "ppt_expected_editable_text_missing"
                                        ),
                                        "page": page,
                                        "field": field,
                                    }
                                )
                    raw_required = requirements.get(str(page), ())
                    required = (
                        tuple(str(item) for item in raw_required)
                        if isinstance(raw_required, Sequence)
                        and not isinstance(raw_required, (str, bytes))
                        else ()
                    )
                    if "table" in required and tables < 1:
                        failures.append(
                            {"code": "ppt_required_table_missing", "page": page}
                        )
                    chart_requirements = [
                        item for item in required if item == "chart" or item.startswith("chart:")
                    ]
                    if chart_requirements and charts < 1:
                        failures.append(
                            {"code": "ppt_required_chart_missing", "page": page}
                        )
                    for requirement in chart_requirements:
                        if ":" not in requirement or not chart_types:
                            continue
                        expected_type = requirement.split(":", 1)[1]
                        matches = {
                            "bar": ("bar", "column"),
                            "line": ("line",),
                            "pie": ("pie",),
                        }.get(expected_type, (expected_type,))
                        if not any(
                            token in chart_type
                            for chart_type in chart_types
                            for token in matches
                        ):
                            failures.append(
                                {
                                    "code": "ppt_required_chart_type_mismatch",
                                    "page": page,
                                    "expected": expected_type,
                                    "actual": chart_types,
                                }
                            )
            except Exception as exc:  # library boundary must fail closed
                failures.append(
                    {
                        "code": "ppt_openxml_invalid",
                        "detail": exc.__class__.__name__,
                    }
                )
    if (
        actual_slide_count is not None
        and actual_slide_count != expected_slide_count
    ):
        failures.append(
            {
                "code": "ppt_slide_count_mismatch",
                "expected": expected_slide_count,
                "actual": actual_slide_count,
            }
        )
    result: dict[str, JsonValue] = {
        "passed": not failures,
        "expected_slide_count": expected_slide_count,
        "actual_slide_count": actual_slide_count,
        "editable_required": editable_required,
        "shape_inventory": shape_inventory,
        "hard_failures": failures,
    }
    validate_json_value(result)
    return result


def _preview_quality_gate(
    previews: Sequence[Mapping[str, Any]],
    *,
    expected_slide_count: int,
    expected_render_hash: str | None = None,
) -> dict[str, JsonValue]:
    pages: list[int] = []
    invalid_pages: list[int] = []
    integrity_failures: list[dict[str, JsonValue]] = []
    for index, preview in enumerate(previews, start=1):
        page = _strict_page_number(preview.get("page"))
        page = page if page is not None else -1
        path = str(preview.get("path") or "")
        candidate = Path(path).expanduser() if path else None
        try:
            valid_file = bool(
                candidate
                and candidate.is_file()
                and candidate.stat().st_size > 0
            )
        except OSError:
            valid_file = False
        if (
            page < 1
            or not valid_file
        ):
            invalid_pages.append(page)
        if valid_file and candidate is not None:
            actual_sha256 = file_or_value_hash(str(candidate), {})
            declared_sha256 = str(preview.get("sha256") or "")
            if (
                declared_sha256 != actual_sha256
                or str(preview.get("artifact_ref") or "")
                != f"preview:{actual_sha256}"
            ):
                integrity_failures.append(
                    {
                        "code": "ppt_preview_hash_mismatch",
                        "page": page,
                        "declared": declared_sha256 or None,
                        "actual": actual_sha256,
                    }
                )
            if (
                expected_render_hash is not None
                and str(preview.get("render_hash") or "") != expected_render_hash
            ):
                integrity_failures.append(
                    {
                        "code": "ppt_preview_render_hash_mismatch",
                        "page": page,
                    }
                )
        pages.append(page)
    expected_pages = list(range(1, expected_slide_count + 1))
    failures: list[dict[str, JsonValue]] = []
    if pages != expected_pages or invalid_pages:
        failures.append(
            {
                "code": "ppt_preview_coverage_incomplete",
                "expected_pages": expected_pages,
                "actual_pages": pages,
                "invalid_pages": invalid_pages,
            }
        )
    failures.extend(integrity_failures)
    result: dict[str, JsonValue] = {
        "passed": not failures,
        "hard_failures": failures,
        "expected_pages": expected_pages,
        "actual_pages": pages,
    }
    validate_json_value(result)
    return result


def _strict_page_number(value: Any) -> int | None:
    if isinstance(value, bool) or not isinstance(value, int) or value < 1:
        return None
    return value


def _review_quality_gate(
    review: Mapping[str, Any],
    *,
    expected_slide_count: int,
    allow_issues: bool = False,
) -> dict[str, JsonValue]:
    failures: list[dict[str, JsonValue]] = []
    raw_issues = review.get("issues")
    issues_valid = isinstance(raw_issues, list) and all(
        isinstance(item, Mapping) for item in raw_issues
    )
    if not issues_valid:
        failures.append({"code": "ppt_visual_review_issues_invalid"})
        issues: list[Mapping[str, Any]] = []
    else:
        issues = raw_issues
        if issues and not allow_issues:
            failures.append(
                {
                    "code": "ppt_visual_issues_remaining",
                    "count": len(issues),
                }
            )

    raw_reviews = review.get("reviews")
    review_pages: list[int] = []
    bad_review_pages: list[int] = []
    reviews_valid = isinstance(raw_reviews, list)
    if reviews_valid:
        for item in raw_reviews:
            if not isinstance(item, Mapping):
                reviews_valid = False
                continue
            page = _strict_page_number(item.get("page"))
            ok = item.get("ok")
            if page is None or not isinstance(ok, bool):
                reviews_valid = False
                continue
            review_pages.append(page)
            if not ok:
                bad_review_pages.append(page)
    if not reviews_valid:
        failures.append({"code": "ppt_visual_reviews_invalid"})

    reviewed_count = _strict_page_number(review.get("reviewed_pages"))
    expected_pages = list(range(1, expected_slide_count + 1))
    if reviewed_count != expected_slide_count or review_pages != expected_pages:
        failures.append(
            {
                "code": "ppt_visual_review_coverage_incomplete",
                "expected_pages": expected_pages,
                "actual_pages": review_pages,
                "reviewed_pages": reviewed_count,
            }
        )

    issue_pages: list[int] = []
    for issue in issues:
        page = _strict_page_number(issue.get("page"))
        if page is None:
            failures.append({"code": "ppt_visual_review_issue_page_invalid"})
        else:
            issue_pages.append(page)
    if sorted(issue_pages) != sorted(bad_review_pages):
        failures.append(
            {
                "code": "ppt_visual_review_result_inconsistent",
                "issue_pages": sorted(issue_pages),
                "failed_review_pages": sorted(bad_review_pages),
            }
        )

    raw_hard_failures = review.get("hard_failures")
    if "hard_failures" in review:
        if isinstance(raw_hard_failures, list):
            failures.extend(
                dict(item) for item in raw_hard_failures if isinstance(item, Mapping)
            )
            if any(not isinstance(item, Mapping) for item in raw_hard_failures):
                failures.append(
                    {"code": "ppt_visual_review_hard_failures_invalid"}
                )
        else:
            failures.append({"code": "ppt_visual_review_hard_failures_invalid"})

    result: dict[str, JsonValue] = {
        "passed": not failures,
        "hard_failures": failures,
        "expected_pages": expected_pages,
        "actual_pages": review_pages,
    }
    validate_json_value(result)
    return result


def validate_delivery_quality(
    *,
    render_ref: Mapping[str, Any],
    previews: Sequence[Mapping[str, Any]],
    review: Mapping[str, Any],
    expected_slide_count: int,
    expected_preview_hash: str,
    editable_required: bool = False,
    slide_requirements: Mapping[str, Sequence[str]] | None = None,
    expected_slides: Sequence[Mapping[str, Any]] | None = None,
) -> dict[str, JsonValue]:
    structural = inspect_rendered_deck(
        str(render_ref.get("path") or "") or None,
        expected_slide_count=expected_slide_count,
        editable_required=editable_required,
        slide_requirements=slide_requirements,
        expected_slides=expected_slides,
    )
    preview = _preview_quality_gate(
        previews,
        expected_slide_count=expected_slide_count,
        expected_render_hash=str(render_ref.get("output_hash") or ""),
    )
    visual = _review_quality_gate(
        review, expected_slide_count=expected_slide_count
    )
    failures = [
        *structural["hard_failures"],
        *preview["hard_failures"],
        *visual["hard_failures"],
    ]
    path = str(render_ref.get("path") or "")
    recorded_hash = str(render_ref.get("output_hash") or "")
    actual_hash = file_or_value_hash(path or None, {})
    if not recorded_hash or actual_hash != recorded_hash:
        failures.append(
            {
                "code": "ppt_output_hash_mismatch",
                "recorded": recorded_hash or None,
                "actual": actual_hash,
            }
        )
    actual_preview_hash = content_hash(
        [dict(item) for item in previews]
    )
    if not expected_preview_hash or actual_preview_hash != expected_preview_hash:
        failures.append(
            {
                "code": "ppt_preview_set_hash_mismatch",
                "recorded": expected_preview_hash or None,
                "actual": actual_preview_hash,
            }
        )
    review_payload = dict(review)
    recorded_review_hash = str(review_payload.pop("review_hash", "") or "")
    actual_review_hash = content_hash(review_payload)
    if not recorded_review_hash or actual_review_hash != recorded_review_hash:
        failures.append(
            {
                "code": "ppt_visual_review_hash_mismatch",
                "recorded": recorded_review_hash or None,
                "actual": actual_review_hash,
            }
        )
    result: dict[str, JsonValue] = {
        "passed": not failures,
        "hard_failures": failures,
        "structural": structural,
        "preview": preview,
        "visual": visual,
    }
    validate_json_value(result)
    return result


def slide_payload(slide: ppt_tools.SlideOutline | Mapping[str, Any]) -> dict[str, JsonValue]:
    parsed = (
        copy.deepcopy(slide).normalize()
        if isinstance(slide, ppt_tools.SlideOutline)
        else ppt_tools.parse_outline([dict(slide)])[0]
    )
    payload = asdict(parsed)
    validate_json_value(payload)
    return payload


def stable_slide_records(
    slides: Sequence[ppt_tools.SlideOutline | Mapping[str, Any]],
    *,
    full_page_images: bool = False,
) -> list[dict[str, JsonValue]]:
    payloads = [slide_payload(slide) for slide in slides]
    planned_payloads = payloads
    if full_page_images:
        planned = ppt_tools.plan_full_page_layouts(
            [ppt_tools.parse_outline([payload])[0] for payload in payloads]
        )
        planned_payloads = [slide_payload(slide) for slide in planned]
    records: list[dict[str, JsonValue]] = []
    for index, payload in enumerate(planned_payloads):
        identity_payload = copy.deepcopy(payloads[index])
        identity_payload["image_path"] = None
        if full_page_images:
            identity_payload["notes"] = ""
            identity_payload["full_page_layout"] = ""
        stable_id = f"slide-{index + 1:03d}-{content_hash(identity_payload)[:12]}"
        if full_page_images:
            prompt, size = _full_page_prompt_payload(payload)
            model = ppt_tools._resolve_ppt_image_model()
            image: dict[str, JsonValue] = {
                "mode": FULL_PAGE_RENDER_MODE,
                "status": "pending",
                "prompt": prompt,
                "size": size,
                "model": model,
                "prompt_schema_version": ppt_tools.FULL_PAGE_PROMPT_SCHEMA_VERSION,
                "normalizer_version": ppt_tools.FULL_PAGE_NORMALIZER_VERSION,
                "layout_spec_version": ppt_tools.FULL_PAGE_LAYOUT_SPEC_VERSION,
                "compositor_version": ppt_tools.FULL_PAGE_COMPOSITOR_VERSION,
                "font_policy_version": ppt_tools.FULL_PAGE_FONT_POLICY_VERSION,
                "layout": str(payload.get("full_page_layout") or "text_left"),
                "page_revision": 0,
                "revision_codes": [],
                "pre_hash": _full_page_pre_hash(
                    slide_id=stable_id,
                    prompt=prompt,
                    size=size,
                    model=model,
                    visible_copy=visible_text_payload(payload),
                    layout=str(payload.get("full_page_layout") or "text_left"),
                    page_revision=0,
                    revision_codes=[],
                ),
                "post_hash": None,
                "path": None,
                "error": None,
            }
        else:
            prompt, size = _prompt_payload(payload)
            image = {
                "status": "pending" if prompt else "not_requested",
                "prompt": prompt,
                "size": size,
                "pre_hash": content_hash(
                    {"slide_id": stable_id, "prompt": prompt, "size": size}
                ),
                "post_hash": None,
                "path": payload.get("image_path"),
                "error": None,
            }
        records.append(
            {
                "slide_id": stable_id,
                "index": index,
                "slide": payload,
                "image": image,
            }
        )
    return records


def _refresh_full_page_image(
    record: dict[str, Any], *, revision: int | None = None, codes: Sequence[str] | None = None
) -> None:
    slide = dict(record["slide"])
    image = dict(record.get("image", {}))
    page_revision = int(image.get("page_revision") or 0) if revision is None else revision
    revision_codes = list(image.get("revision_codes", [])) if codes is None else list(codes)
    prompt, size = _full_page_prompt_payload(
        slide, page_revision=page_revision, revision_codes=revision_codes
    )
    model = str(image.get("model") or ppt_tools._resolve_ppt_image_model())
    layout = str(slide.get("full_page_layout") or "text_left")
    image.update(
        {
            "mode": FULL_PAGE_RENDER_MODE,
            "status": "pending",
            "prompt": prompt,
            "size": size,
            "model": model,
            "prompt_schema_version": ppt_tools.FULL_PAGE_PROMPT_SCHEMA_VERSION,
            "normalizer_version": ppt_tools.FULL_PAGE_NORMALIZER_VERSION,
            "layout_spec_version": ppt_tools.FULL_PAGE_LAYOUT_SPEC_VERSION,
            "compositor_version": ppt_tools.FULL_PAGE_COMPOSITOR_VERSION,
            "font_policy_version": ppt_tools.FULL_PAGE_FONT_POLICY_VERSION,
            "layout": layout,
            "page_revision": page_revision,
            "revision_codes": sorted(revision_codes),
            "pre_hash": _full_page_pre_hash(
                slide_id=str(record["slide_id"]),
                prompt=prompt,
                size=size,
                model=model,
                visible_copy=visible_text_payload(slide),
                layout=layout,
                page_revision=page_revision,
                revision_codes=revision_codes,
            ),
            "post_hash": None,
            "path": None,
            "error": None,
        }
    )
    slide["image_path"] = None
    record["slide"] = slide
    record["image"] = image


def normalize_full_page_records(
    records: Sequence[Mapping[str, Any]],
) -> list[dict[str, JsonValue]]:
    """Upgrade legacy durable records without repeating committed image effects."""
    copied = copy.deepcopy(list(records))
    full_page = [
        record for record in copied
        if isinstance(record.get("image"), Mapping)
        and record["image"].get("mode") == FULL_PAGE_RENDER_MODE
    ]
    needs_upgrade = any(
        str(record.get("slide", {}).get("full_page_layout") or "") not in ppt_tools.FULL_PAGE_LAYOUTS
        or (
            record.get("image", {}).get("status") == "ready"
            and not Path(
                str(
                    record.get("image", {}).get("path")
                    or record.get("slide", {}).get("image_path")
                    or ""
                )
            ).is_file()
        )
        for record in full_page
    )
    if not full_page or not needs_upgrade:
        return copied  # type: ignore[return-value]

    ordered = sorted(copied, key=lambda item: int(item.get("index", 0)))
    slides = [ppt_tools.parse_outline([dict(record["slide"])])[0] for record in ordered]
    locked: dict[int, str] = {}
    ready_ids: set[str] = set()
    for position, record in enumerate(ordered):
        image = dict(record.get("image", {}))
        path = str(image.get("path") or record.get("slide", {}).get("image_path") or "")
        if image.get("status") == "ready" and path and Path(path).is_file():
            ready_ids.add(str(record.get("slide_id") or ""))
            locked[position] = str(record["slide"].get("full_page_layout") or "text_left")
    planned = ppt_tools.plan_full_page_layouts(
        slides, locked_layouts=locked, strict_coverage=False
    )
    for position, (record, slide) in enumerate(zip(ordered, planned)):
        slide_id = str(record.get("slide_id") or "")
        payload = slide_payload(slide)
        if slide_id in ready_ids:
            existing = dict(record["slide"])
            existing["full_page_layout"] = locked[position]
            record["slide"] = existing
            continue
        record["slide"] = payload
        _refresh_full_page_image(record)
    validate_json_value(copied)
    return copied  # type: ignore[return-value]


def records_to_slides(records: Sequence[Mapping[str, Any]]) -> list[ppt_tools.SlideOutline]:
    ordered = sorted(records, key=lambda item: int(item.get("index", 0)))
    return [ppt_tools.parse_outline([dict(item["slide"])])[0] for item in ordered]


def prepare_slide_records(
    records: Sequence[Mapping[str, Any]], *, image_mode: bool, reachable: bool
) -> tuple[list[dict[str, JsonValue]], str]:
    copied = copy.deepcopy(list(records))
    is_full_page = bool(copied) and all(
        isinstance(record.get("image"), Mapping)
        and record["image"].get("mode") == FULL_PAGE_RENDER_MODE
        for record in copied
    )
    render_mode = (
        FULL_PAGE_RENDER_MODE
        if image_mode and reachable and is_full_page
        else "images" if image_mode and reachable else "template"
    )
    if render_mode == "template":
        degraded = ppt_tools._degrade_to_template(records_to_slides(copied))
        for record, slide in zip(copied, degraded):
            record["slide"] = slide_payload(slide)
            image = dict(record.get("image", {}))
            if image.get("status") == "pending":
                image["status"] = "skipped_template"
            record["image"] = image
    return copied, render_mode


def next_pending_slide(records: Sequence[Mapping[str, Any]]) -> dict[str, JsonValue] | None:
    for record in sorted(records, key=lambda item: int(item.get("index", 0))):
        image = record.get("image")
        if isinstance(image, Mapping) and image.get("status") == "pending":
            return copy.deepcopy(dict(record))  # type: ignore[return-value]
    return None


async def probe_images(port: object | None, *, timeout_s: float) -> bool:
    method = _port_method(port, "probe_images", "probe_image_reachable", "probe")
    if method is None:
        return bool(ppt_tools.probe_image_reachable(timeout_s=timeout_s))
    return bool(await _await(method(timeout_s=timeout_s)))


async def generate_slide_image(
    record: Mapping[str, Any], port: object | None
) -> dict[str, JsonValue]:
    updated = copy.deepcopy(dict(record))
    image = dict(updated.get("image", {}))
    prompt = str(image.get("prompt") or "")
    size = str(image.get("size") or "1536x1024")
    full_page = image.get("mode") == FULL_PAGE_RENDER_MODE
    model = str(image.get("model") or "") or None
    method = _port_method(port, "generate_slide_image", "generate_image")
    if method is None:
        kwargs: dict[str, Any] = {"size": size}
        if model:
            kwargs["model"] = model
        raw = ppt_tools.generate_images([prompt], **kwargs)
        result = raw[0] if raw else {}
    else:
        call_kwargs: dict[str, Any] = {
            "slide_id": str(updated["slide_id"]),
            "prompt": prompt,
            "size": size,
            "input_hash": str(image["pre_hash"]),
        }
        if full_page:
            call_kwargs.update(
                {
                    "model": model,
                    "prompt_schema_version": str(image["prompt_schema_version"]),
                    "normalizer_version": str(image["normalizer_version"]),
                    "page_revision": int(image.get("page_revision") or 0),
                    "full_page": True,
                    "slide": copy.deepcopy(updated["slide"]),
                }
            )
        result = await _await(method(**call_kwargs))
    result = dict(result or {}) if isinstance(result, Mapping) else {"error": str(result)}
    path = str(result.get("path") or "") or None
    if full_page and path and Path(path).is_file() and not ppt_tools._is_normalized_full_page_image(path):
        try:
            path = ppt_tools._compose_full_page_image(path, updated["slide"])
            result["path"] = path
        except Exception as exc:  # noqa: BLE001
            result = {
                "path": None,
                "error": f"full-page normalization failed: {exc}",
                "error_kind": "normalization_failed",
            }
            path = None
    if path and Path(path).is_file():
        image.update(
            {
                "status": "ready",
                "path": path,
                "post_hash": file_or_value_hash(path, result),
                "error": None,
            }
        )
        slide = dict(updated["slide"])
        slide["image_path"] = path
        updated["slide"] = slide
    else:
        error_kind = str(
            result.get("error_kind")
            or ("missing_page_image" if path else "generation_failed")
        )
        image.update(
            {
                "status": "fallback_required"
                if error_kind in {"connectivity", "model_unavailable"}
                else "failed",
                "post_hash": content_hash(result),
                "error": str(result.get("error") or error_kind),
            }
        )
    updated["image"] = image
    validate_json_value(updated)
    return updated  # type: ignore[return-value]


async def render_deck(
    records: Sequence[Mapping[str, Any]],
    *,
    port: object | None,
    topic: str,
    title: str,
    author: str,
    theme: str,
    output_path: str | None,
    render_mode: str,
    render_revision: int,
    editable_required: bool = False,
    slide_requirements: Mapping[str, Sequence[str]] | None = None,
) -> dict[str, JsonValue]:
    slides = records_to_slides(records)
    input_payload: dict[str, JsonValue] = {
        "slides": [dict(record) for record in records],
        "topic": topic,
        "title": title,
        "author": author,
        "theme": theme,
        "output_path": output_path,
        "render_mode": render_mode,
        "render_revision": render_revision,
        "editable_required": editable_required,
        "slide_requirements": copy.deepcopy(dict(slide_requirements or {})),
    }
    input_hash = content_hash(input_payload)
    method = _port_method(port, "render_ppt", "render_deck", "render")
    if method is None:
        if render_mode == FULL_PAGE_RENDER_MODE:
            result = ppt_tools._render_full_page_images(
                slides,
                title=title,
                author=author,
                output_path=output_path,
            )
        else:
            result = ppt_tools._render_pro(
                slides,
                theme=theme,
                title=title,
                author=author,
                output_path=output_path,
                image_mode=render_mode == "images",
                probe_timeout_s=0.1,
                notify=lambda _message: None,
                editable_required=editable_required,
                images_prepared=render_mode == "images",
            )
    else:
        result = await _await(
            method(
                slides=[asdict(slide) for slide in slides],
                topic=topic,
                title=title,
                author=author,
                theme=theme,
                output_path=output_path,
                render_mode=render_mode,
                render_revision=render_revision,
                input_hash=input_hash,
                editable_required=editable_required,
            )
        )
    payload = dict(result or {}) if isinstance(result, Mapping) else {"ok": False, "error": str(result)}
    payload["input_hash"] = input_hash
    if bool(payload.get("ok")):
        structural_gate = inspect_rendered_deck(
            str(payload.get("path") or "") or None,
            expected_slide_count=len(records),
            editable_required=editable_required,
            slide_requirements=slide_requirements,
            expected_slides=records,
        )
        payload["structural_quality_gate"] = structural_gate
        if not structural_gate["passed"]:
            payload.update(
                {
                    "ok": False,
                    "error_code": "ppt_structural_quality_failed",
                    "error": "PPT file structure or slide count failed validation",
                }
            )
    path = str(payload.get("path") or "") or None
    payload["output_hash"] = file_or_value_hash(path, payload)
    validate_json_value(payload)
    return payload  # type: ignore[return-value]


async def render_preview(render_ref: Mapping[str, Any], port: object | None) -> list[dict[str, JsonValue]]:
    page_images = render_ref.get("page_images", [])
    if (
        render_ref.get("render_mode") == FULL_PAGE_RENDER_MODE
        and isinstance(page_images, Sequence)
        and not isinstance(page_images, (str, bytes))
    ):
        refs: list[dict[str, JsonValue]] = []
        render_hash = str(render_ref.get("output_hash") or "")
        for index, item in enumerate(page_images, start=1):
            if not isinstance(item, Mapping):
                continue
            path = str(item.get("path") or "")
            if not path or not Path(path).is_file():
                continue
            page = _strict_page_number(item.get("page"))
            if page is None:
                continue
            sha256 = file_or_value_hash(path, {"page": page, "path": path})
            refs.append(
                {
                    "kind": "image",
                    "path": path,
                    "mime": "image/png",
                    "title": f"Preview slide {page}",
                    "page": page,
                    "sha256": sha256,
                    "artifact_ref": f"preview:{sha256}",
                    "render_hash": render_hash,
                }
            )
        validate_json_value(refs)
        return refs
    method = _port_method(port, "render_preview", "preview_ppt", "preview")
    if method is None:
        artifacts = render_ref.get("artifacts", [])
        return [dict(item) for item in artifacts if isinstance(item, Mapping) and item.get("kind") == "image"]
    raw = await _await(
        method(
            path=str(render_ref.get("path") or ""),
            render_hash=str(render_ref.get("output_hash") or ""),
        )
    )
    values = raw if isinstance(raw, Sequence) and not isinstance(raw, (str, bytes)) else []
    result = [dict(item) for item in values if isinstance(item, Mapping)]
    validate_json_value(result)
    return result  # type: ignore[return-value]


async def evaluate_visuals(
    previews: Sequence[Mapping[str, Any]],
    records: Sequence[Mapping[str, Any]],
    evaluator: object | None,
) -> dict[str, JsonValue]:
    expected_slide_count = len(records)
    preview_gate = _preview_quality_gate(
        previews, expected_slide_count=expected_slide_count
    )
    if not preview_gate["passed"]:
        result: dict[str, JsonValue] = {
            "issues": [],
            "score": None,
            "reviews": [],
            "reviewed_pages": 0,
            "preview_quality_gate": preview_gate,
            "hard_failures": list(preview_gate["hard_failures"]),
        }
        result["review_hash"] = content_hash(result)
        validate_json_value(result)
        return result
    method = _port_method(evaluator, "evaluate_ppt", "review_slides", "evaluate")
    if method is None:
        result = {
            "issues": [],
            "score": None,
            "reviews": [],
            "reviewed_pages": 0,
            "preview_quality_gate": preview_gate,
            "hard_failures": [{"code": "ppt_visual_evaluator_unavailable"}],
        }
        result["review_hash"] = content_hash(result)
        validate_json_value(result)
        return result
    raw = await _await(
        method(
            previews=[dict(item) for item in previews],
            slides=[dict(item) for item in records],
        )
    )
    if isinstance(raw, Mapping):
        result = dict(raw)
    else:
        result = {"issues": list(raw or [])}
    issues = result.get("issues")
    if not isinstance(issues, list) or not all(
        isinstance(item, Mapping) for item in issues
    ):
        result["issues"] = []
        result["hard_failures"] = [
            {"code": "ppt_visual_review_issues_invalid"}
        ]
    else:
        result["issues"] = issues
    reviews = result.get("reviews", [])
    review_pages: list[int] = []
    if isinstance(reviews, Sequence) and not isinstance(reviews, (str, bytes)):
        for review in reviews:
            if not isinstance(review, Mapping):
                continue
            page = _strict_page_number(review.get("page"))
            review_pages.append(page if page is not None else 0)
    reviewed_pages = result.get("reviewed_pages")
    reviewed_count = _strict_page_number(reviewed_pages)
    raw_hard_failures = result.get("hard_failures")
    hard_failures: list[dict[str, Any]] = []
    if isinstance(raw_hard_failures, list):
        hard_failures.extend(
            dict(item) for item in raw_hard_failures if isinstance(item, Mapping)
        )
        if any(not isinstance(item, Mapping) for item in raw_hard_failures):
            hard_failures.append(
                {"code": "ppt_visual_review_hard_failures_invalid"}
            )
    elif "hard_failures" in result:
        hard_failures.append({"code": "ppt_visual_review_hard_failures_invalid"})
    result["reviews"] = list(reviews) if isinstance(reviews, list) else []
    result["reviewed_pages"] = reviewed_count or 0
    result["preview_quality_gate"] = preview_gate
    result["hard_failures"] = hard_failures
    review_gate = _review_quality_gate(
        result,
        expected_slide_count=expected_slide_count,
        allow_issues=True,
    )
    result["hard_failures"] = list(review_gate["hard_failures"])
    result["review_hash"] = content_hash(result)
    validate_json_value(result)
    return result  # type: ignore[return-value]


def apply_visual_revision(
    records: Sequence[Mapping[str, Any]], issues: Sequence[Mapping[str, Any]]
) -> list[dict[str, JsonValue]]:
    copied = copy.deepcopy(list(records))
    by_index = {int(item.get("index", -1)): item for item in copied}
    by_id = {str(item.get("slide_id", "")): item for item in copied}
    grouped: dict[str, set[str]] = {}
    for issue in issues:
        record = by_id.get(str(issue.get("slide_id") or ""))
        if record is None:
            raw_page = issue.get("page", issue.get("slide_index", -1))
            try:
                page = int(raw_page)
            except (TypeError, ValueError):
                continue
            record = by_index.get(page - 1 if page > 0 else page)
        if record is None:
            continue
        image = record.get("image", {})
        if isinstance(image, Mapping) and image.get("mode") == FULL_PAGE_RENDER_MODE:
            slide_id = str(record.get("slide_id") or "")
            grouped.setdefault(slide_id, set()).update(_normalized_revision_codes(issue))
            continue
        slide = dict(record["slide"])
        action = str(issue.get("action") or "shrink_text")
        if action == "change_variant" and issue.get("variant"):
            variant = str(issue["variant"]).strip().lower()
            if (
                str(slide.get("layout") or "") in {"chart", "toc"}
                and variant not in ppt_tools.IMAGE_VARIANTS
            ):
                variant = "top"
            slide["image_variant"] = variant
        elif action == "change_layout" and issue.get("layout"):
            slide["layout"] = str(issue["layout"])
        else:
            current = float(slide.get("font_scale") or 1.0)
            slide["font_scale"] = max(0.5, round(current * 0.9, 3))
        record["slide"] = slide

    for slide_id, code_set in grouped.items():
        record = by_id[slide_id]
        slide = dict(record["slide"])
        image = dict(record["image"])
        revision = int(image.get("page_revision") or 0) + 1
        codes = sorted(code_set or {"text_unreadable"})
        if "layout_misfit" in codes:
            index = int(record.get("index", 0))
            neighbors = {
                str(item.get("slide", {}).get("full_page_layout") or "")
                for item in copied
                if abs(int(item.get("index", -99)) - index) == 1
            }
            slide["full_page_layout"] = ppt_tools.next_full_page_layout(
                slide,
                str(slide.get("full_page_layout") or image.get("layout") or "text_left"),
                forbidden=neighbors,
            )
        record["slide"] = slide
        record["image"] = image
        _refresh_full_page_image(record, revision=revision, codes=codes)
    validate_json_value(copied)
    return copied  # type: ignore[return-value]


def visible_text_payload(slide: Mapping[str, Any]) -> dict[str, JsonValue]:
    parsed = ppt_tools.parse_outline([dict(slide)])[0]
    return {
        "title": parsed.title,
        "subtitle": parsed.subtitle,
        "bullets": list(parsed.bullets),
        "left_title": parsed.left_title,
        "left": list(parsed.left),
        "right_title": parsed.right_title,
        "right": list(parsed.right),
        "quote": parsed.quote,
        "cite": parsed.cite,
        "caption": parsed.caption,
    }


def _full_page_prompt_payload(
    slide: Mapping[str, Any],
    *,
    page_revision: int = 0,
    revision_codes: Sequence[str] = (),
) -> tuple[str, str]:
    parsed = ppt_tools.parse_outline([dict(slide)])[0]
    visual_direction = str(parsed.image_prompt or "").strip()
    visual_direction = re.sub(
        r"\bno\s+(?:text|words?|letters?|typography)\b[,; ]*",
        "",
        visual_direction,
        flags=re.IGNORECASE,
    ).strip()
    safe_codes = sorted({str(code) for code in revision_codes if code in _FULL_PAGE_REVISION_CODES})
    revision_text = " ".join(
        _REVISION_INSTRUCTIONS[code]
        for code in safe_codes
        if code == "layout_misfit"
    )
    layout = parsed.full_page_layout or "text_left"
    spec = ppt_tools.FULL_PAGE_LAYOUT_SPECS[layout]
    prompt = (
        "Create a professional 16:9 presentation visual background only. "
        "The application adds all typography after generation. "
        f"LAYOUT={layout}. {spec.negative_space_instruction} "
        "Render absolutely no words, letters, numbers, labels, UI text, logos, "
        "signatures, or watermarks anywhere in the image. "
        f"VISUAL_DIRECTION={visual_direction or 'clean editorial presentation design'}. "
        f"PAGE_REVISION={int(page_revision)}. {revision_text}"
    ).strip()
    return prompt, ppt_tools.FULL_PAGE_IMAGE_SIZE


def _full_page_pre_hash(
    *,
    slide_id: str,
    prompt: str,
    size: str,
    model: str,
    visible_copy: Mapping[str, JsonValue],
    layout: str = "text_left",
    page_revision: int,
    revision_codes: Sequence[str],
) -> str:
    return content_hash(
        {
            "mode": FULL_PAGE_RENDER_MODE,
            "slide_id": slide_id,
            "prompt": prompt,
            "size": size,
            "model": model,
            "visible_copy": dict(visible_copy),
            "layout": layout,
            "layout_spec_version": ppt_tools.FULL_PAGE_LAYOUT_SPEC_VERSION,
            "prompt_schema_version": ppt_tools.FULL_PAGE_PROMPT_SCHEMA_VERSION,
            "normalizer_version": ppt_tools.FULL_PAGE_NORMALIZER_VERSION,
            "compositor_version": ppt_tools.FULL_PAGE_COMPOSITOR_VERSION,
            "font_policy_version": ppt_tools.FULL_PAGE_FONT_POLICY_VERSION,
            "target_size": list(ppt_tools.FULL_PAGE_TARGET_SIZE),
            "page_revision": page_revision,
            "revision_codes": list(sorted(revision_codes)),
        }
    )


def _normalized_revision_codes(issue: Mapping[str, Any]) -> set[str]:
    raw_codes = issue.get("reason_codes", [])
    if not isinstance(raw_codes, Sequence) or isinstance(raw_codes, (str, bytes)):
        raw_codes = []
    codes = {
        str(code)
        for code in raw_codes
        if str(code) in _FULL_PAGE_REVISION_CODES
    }
    action = str(issue.get("action") or "")
    action_map = {
        "change_layout": "layout_misfit",
        "change_variant": "layout_misfit",
        "regenerate_page": "text_incorrect",
        "shrink_text": "text_overflow",
    }
    if action in action_map:
        codes.add(action_map[action])
    return codes or {"text_unreadable"}


def _prompt_payload(slide: Mapping[str, Any]) -> tuple[str, str]:
    parsed = ppt_tools.parse_outline([dict(slide)])[0]
    if not parsed.image_prompt:
        return "", ""
    return ppt_tools._image_prompt_payload(parsed)


def _port_method(port: object | None, *names: str):
    if port is None:
        return None
    for name in names:
        method = getattr(port, name, None)
        if callable(method):
            return method
    return None


async def _await(value: Any) -> Any:
    return await value if inspect.isawaitable(value) else value


__all__ = [
    "FULL_PAGE_RENDER_MODE",
    "apply_visual_revision",
    "content_hash",
    "evaluate_visuals",
    "file_or_value_hash",
    "generate_slide_image",
    "next_pending_slide",
    "prepare_slide_records",
    "normalize_full_page_records",
    "probe_images",
    "records_to_slides",
    "render_deck",
    "render_preview",
    "slide_payload",
    "stable_slide_records",
    "visible_text_payload",
]
