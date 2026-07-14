from __future__ import annotations

import hashlib
from pathlib import Path

import pytest

from deskpet.tools import ppt_tools
from deskpet.workflows import WorkflowContext
from deskpet.workflows.adapters.ppt_runtime import PptRuntime
from deskpet.workflows.definitions import ppt_pro_nodes as nodes
from deskpet.workflows.definitions.v1 import ppt_pro as ppt_pro_graph


pytestmark = pytest.mark.skipif(
    not ppt_tools._HAS_PPTX, reason="python-pptx not installed"
)


def _provider_image(path: Path, *, color: tuple[int, int, int] = (20, 40, 80)) -> Path:
    from PIL import Image

    Image.new("RGB", (1792, 1024), color).save(path)
    return path


def _page_image(path: Path, *, color: tuple[int, int, int] = (20, 40, 80)) -> Path:
    return Path(
        ppt_tools._normalize_full_page_image(
            str(_provider_image(path, color=color))
        )
    )


def test_full_page_records_keep_exact_copy_out_of_background_prompt_and_in_hash(monkeypatch) -> None:
    monkeypatch.setattr(ppt_tools, "_resolve_ppt_image_model", lambda: "model-a")
    slide = {
        "layout": "bullet",
        "title": "精确标题 2026",
        "subtitle": "副标题",
        "bullets": ["第一条", "Second item"],
        "left": ["左侧"],
        "right": ["右侧"],
        "quote": "引用原文",
        "cite": "来源",
        "caption": "图注",
        "notes": "绝不能进入图片的演讲备注",
    }

    first = nodes.stable_slide_records([slide], full_page_images=True)[0]
    prompt = first["image"]["prompt"]
    assert first["image"]["status"] == "pending"
    assert first["image"]["size"] == "1792x1024"
    assert first["image"]["model"] == "model-a"
    for text in ("精确标题 2026", "副标题", "第一条", "Second item", "左侧", "右侧", "引用原文", "来源", "图注"):
        assert text not in prompt
    assert "adds all typography after generation" in prompt
    assert "绝不能进入图片的演讲备注" not in prompt

    changed_copy = nodes.stable_slide_records(
        [dict(slide, title="另一个精确标题")], full_page_images=True
    )[0]
    assert changed_copy["image"]["pre_hash"] != first["image"]["pre_hash"]

    changed_notes = dict(slide, notes="另一份备注")
    same_effect = nodes.stable_slide_records(
        [changed_notes], full_page_images=True
    )[0]
    assert same_effect["slide_id"] == first["slide_id"]
    assert same_effect["image"]["pre_hash"] == first["image"]["pre_hash"]

    monkeypatch.setattr(ppt_tools, "_resolve_ppt_image_model", lambda: "model-b")
    changed_model = nodes.stable_slide_records([slide], full_page_images=True)[0]
    assert changed_model["image"]["pre_hash"] != first["image"]["pre_hash"]

    monkeypatch.setattr(ppt_tools, "_resolve_ppt_image_model", lambda: "model-a")
    monkeypatch.setattr(ppt_tools, "FULL_PAGE_PROMPT_SCHEMA_VERSION", "prompt-v2")
    changed_schema = nodes.stable_slide_records([slide], full_page_images=True)[0]
    assert changed_schema["image"]["pre_hash"] != first["image"]["pre_hash"]

    monkeypatch.setattr(
        ppt_tools,
        "FULL_PAGE_PROMPT_SCHEMA_VERSION",
        "ppt-full-page-background-prompt-v4",
    )
    monkeypatch.setattr(ppt_tools, "FULL_PAGE_TARGET_SIZE", (1600, 900))
    changed_target = nodes.stable_slide_records([slide], full_page_images=True)[0]
    assert changed_target["image"]["pre_hash"] != first["image"]["pre_hash"]


def test_cjk_font_missing_fails_closed_without_default_font(monkeypatch) -> None:
    from PIL import ImageFont

    def missing(*args, **kwargs):
        raise OSError("missing")

    monkeypatch.setattr(ImageFont, "truetype", missing)
    with pytest.raises(ppt_tools.FullPageLayoutError) as exc_info:
        ppt_tools._full_page_font(24, text="中文标点，。！？")
    assert exc_info.value.code == "cjk_font_unavailable"


def test_cjk_missing_glyph_rejects_candidate_font(monkeypatch) -> None:
    from PIL import ImageFont

    fallback = ImageFont.load_default()
    monkeypatch.setattr(ImageFont, "truetype", lambda *args, **kwargs: fallback)
    monkeypatch.setattr(
        ppt_tools, "_font_cmap", lambda _path: frozenset({ord("中")})
    )
    with pytest.raises(ppt_tools.FullPageLayoutError) as exc_info:
        ppt_tools._full_page_font(24, text="中𠀀")
    assert exc_info.value.code == "cjk_font_unavailable"


@pytest.mark.parametrize(
    "text",
    [
        "中文标点，不能落在错误位置。Mixed English 2026！",
        "SupercalifragilisticexpialidociousWithoutAnyBreakPoint",
    ],
)
def test_full_page_wrapping_preserves_mixed_copy_and_long_tokens(text: str) -> None:
    from PIL import Image, ImageDraw, ImageFont

    draw = ImageDraw.Draw(Image.new("RGB", (400, 100)))
    lines = ppt_tools._wrap_full_page_text(
        draw, text, ImageFont.load_default(), 90
    )
    assert "".join(lines) == text
    assert all(
        not line.startswith(tuple("，。！？；：、）》】』」〉〕）"))
        for line in lines[1:]
    )


def test_full_page_pre_hash_tracks_normalizer_version(monkeypatch) -> None:
    hash_args = {
        "slide_id": "slide-1",
        "prompt": "full page",
        "size": "1792x1024",
        "model": "model-a",
        "visible_copy": {"title": "Exact title"},
        "page_revision": 0,
        "revision_codes": [],
    }
    baseline = nodes._full_page_pre_hash(**hash_args)

    with monkeypatch.context() as patch:
        patch.setattr(
            ppt_tools,
            "FULL_PAGE_NORMALIZER_VERSION",
            "ppt-full-page-normalizer-audit-v2",
        )
        changed = nodes._full_page_pre_hash(**hash_args)

    assert changed != baseline
    assert nodes._full_page_pre_hash(**hash_args) == baseline


def test_full_page_planner_is_deterministic_and_uses_five_layouts() -> None:
    slides = ppt_tools.parse_outline(
        [
            {"title": f"第 {index + 1} 页", "bullets": ["要点一", "要点二"]}
            for index in range(8)
        ]
    )
    first = ppt_tools.plan_full_page_layouts(slides)
    second = ppt_tools.plan_full_page_layouts(slides)
    layouts = [slide.full_page_layout for slide in first]

    assert layouts == [slide.full_page_layout for slide in second]
    assert len(set(layouts)) >= 5
    assert all(left != right for left, right in zip(layouts, layouts[1:]))
    assert layouts[0] == "cover_band"


def test_six_page_deck_uses_five_layouts_without_adjacent_repeat() -> None:
    slides = ppt_tools.parse_outline(
        [
            {"title": f"第 {index + 1} 页", "bullets": ["要点一", "要点二"]}
            for index in range(6)
        ]
    )

    layouts = [
        slide.full_page_layout for slide in ppt_tools.plan_full_page_layouts(slides)
    ]

    assert len(set(layouts)) >= 5
    assert all(left != right for left, right in zip(layouts, layouts[1:]))


def test_visible_copy_changes_stable_slide_id() -> None:
    first = nodes.stable_slide_records(
        [{"title": "原始标题", "bullets": ["内容"]}], full_page_images=True
    )[0]
    changed = nodes.stable_slide_records(
        [{"title": "更新后的标题", "bullets": ["内容"]}], full_page_images=True
    )[0]

    assert changed["slide_id"] != first["slide_id"]


def test_explicit_full_page_layout_is_a_preference_and_invalid_value_clears() -> None:
    preferred = ppt_tools.parse_outline(
        [{"title": "Preference", "full_page_layout": "text_right"}]
    )[0]
    invalid = ppt_tools.parse_outline(
        [{"title": "Invalid", "full_page_layout": "not-a-layout"}]
    )[0]

    assert ppt_tools.plan_full_page_layouts([preferred])[0].full_page_layout == "text_right"
    assert invalid.full_page_layout == ""


def test_full_page_pre_hash_tracks_layout_and_compositor_versions(monkeypatch) -> None:
    args = {
        "slide_id": "slide-1",
        "prompt": "background",
        "size": "1792x1024",
        "model": "model-a",
        "visible_copy": {"title": "Exact"},
        "layout": "text_left",
        "page_revision": 0,
        "revision_codes": [],
    }
    baseline = nodes._full_page_pre_hash(**args)
    assert nodes._full_page_pre_hash(**dict(args, layout="text_right")) != baseline
    with monkeypatch.context() as patch:
        patch.setattr(ppt_tools, "FULL_PAGE_COMPOSITOR_VERSION", "compositor-next")
        assert nodes._full_page_pre_hash(**args) != baseline


def test_full_page_compositor_flattens_exact_copy_into_safe_page(tmp_path) -> None:
    from PIL import Image, ImageChops

    source = _provider_image(tmp_path / "background.png", color=(80, 100, 130))
    output = Path(
        ppt_tools._compose_full_page_image(
            str(source),
            {
                "title": "确定性中文标题",
                "subtitle": "不会交给图片模型渲染",
                "bullets": ["第一条精确文案", "Second exact item"],
                "cite": "来源：DeskPet",
                "notes": "绝不能显示的演讲备注",
            },
        )
    )

    with Image.open(output) as composed, Image.open(
        ppt_tools._normalize_full_page_image(str(source))
    ) as background:
        assert composed.size == ppt_tools.FULL_PAGE_TARGET_SIZE
        assert ImageChops.difference(composed.convert("RGB"), background.convert("RGB")).getbbox()
        assert composed.getpixel((30, 30)) == background.convert("RGB").getpixel((30, 30))
        assert composed.getpixel((1700, 970)) == composed.getpixel((1700, 850))


def test_full_page_compositor_preserves_uncovered_bottom_content(tmp_path) -> None:
    from PIL import Image

    source = tmp_path / "bottom-detail.png"
    background = Image.new("RGB", ppt_tools.FULL_PAGE_TARGET_SIZE, (20, 40, 80))
    for y in range(880, ppt_tools.FULL_PAGE_TARGET_SIZE[1]):
        color = (y % 256, (y * 2) % 256, (y * 3) % 256)
        for x in range(1500, ppt_tools.FULL_PAGE_TARGET_SIZE[0]):
            background.putpixel((x, y), color)
    background.save(source)

    output = Path(
        ppt_tools._compose_full_page_image(
            str(source),
            {"title": "Deck", "bullets": ["Exact copy"]},
        )
    )

    with Image.open(output) as composed, Image.open(source) as original:
        assert composed.convert("RGB").getpixel((1700, 970)) == original.getpixel(
            (1700, 970)
        )


def test_all_six_full_page_compositors_produce_distinct_safe_pages(tmp_path) -> None:
    source = _provider_image(tmp_path / "six-layouts.png", color=(70, 90, 120))
    digests = set()
    for layout in ppt_tools.FULL_PAGE_LAYOUTS:
        output = Path(
            ppt_tools._compose_full_page_image(
                str(source),
                {
                    "title": "确定性标题",
                    "subtitle": "精确副标题",
                    "bullets": ["第一条", "第二条"],
                    "full_page_layout": layout,
                },
                output_path=str(tmp_path / f"{layout}.png"),
            )
        )
        assert output.is_file()
        digests.add(hashlib.sha256(output.read_bytes()).hexdigest())
    assert len(digests) == len(ppt_tools.FULL_PAGE_LAYOUTS)


def test_full_page_preflight_fails_closed_without_partial_output() -> None:
    with pytest.raises(ppt_tools.FullPageLayoutError) as exc_info:
        ppt_tools.preflight_full_page_copy(
            {"title": "超长内容", "bullets": ["中文内容" * 500] * 20},
            "floating_card",
        )
    assert exc_info.value.code == "text_overflow"


def test_legacy_image_records_keep_images_mode() -> None:
    records = nodes.stable_slide_records(
        [{"layout": "image_full", "title": "Legacy", "image_prompt": "scene, no text"}]
    )
    prepared, render_mode = nodes.prepare_slide_records(
        records, image_mode=True, reachable=True
    )
    assert render_mode == "images"
    assert prepared[0]["image"].get("mode") is None
    assert "no text" in str(prepared[0]["image"]["prompt"]).lower()


@pytest.mark.asyncio
async def test_legacy_state_without_full_page_flag_prepares_legacy_images_mode() -> None:
    from deskpet.workflows import WorkflowContext
    from deskpet.workflows.definitions.v1 import ppt_pro as ppt_graph

    state = ppt_graph.initial_state(topic="Legacy", run_id="legacy-run")
    values = dict(state["values"])
    values.pop("full_page_images", None)
    values.update(
        {
            "outline_slides": [
                {
                    "layout": "image_full",
                    "title": "Legacy",
                    "image_prompt": "legacy scene, no text",
                }
            ],
            "image_probe": {"requested": True, "reachable": True},
        }
    )
    state["values"] = values

    patch = await ppt_graph.prepare_slides_handler(state, WorkflowContext())

    assert patch.values["values"]["render_mode"] == "images"
    assert patch.values["values"]["slide_records"][0]["image"].get("mode") is None


def test_normalizer_and_full_page_assembler_create_one_picture_per_slide(tmp_path: Path) -> None:
    from PIL import Image
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    page1 = _page_image(tmp_path / "raw-1.png")
    page2 = _page_image(tmp_path / "raw-2.png", color=(80, 30, 20))
    with Image.open(page1) as image:
        assert image.size == (1792, 1008)

    output = tmp_path / "full-page.pptx"
    slides = ppt_tools.parse_outline(
        [
            {"title": "One", "image_path": str(page1), "notes": "speaker one"},
            {"title": "Two", "image_path": str(page2), "notes": "speaker two"},
        ]
    )
    result = ppt_tools._render_full_page_images(
        slides, title="Deck", author="DeskPet", output_path=str(output)
    )

    assert result["ok"] is True
    deck = Presentation(output)
    assert len(deck.slides) == 2
    for slide in deck.slides:
        assert len(slide.shapes) == 1
        assert slide.shapes[0].shape_type == MSO_SHAPE_TYPE.PICTURE
        assert not slide.shapes[0].has_text_frame
        assert slide.shapes[0].left == 0 and slide.shapes[0].top == 0
        assert slide.shapes[0].width == deck.slide_width
        assert slide.shapes[0].height == deck.slide_height
    assert "speaker one" in deck.slides[0].notes_slide.notes_text_frame.text
    assert "speaker two" in deck.slides[1].notes_slide.notes_text_frame.text


@pytest.mark.parametrize("source_mode", ["RGBA", "CMYK"])
def test_full_page_normalizer_accepts_common_provider_color_modes(
    tmp_path: Path, source_mode: str
) -> None:
    from PIL import Image

    source = tmp_path / f"provider-{source_mode}.png"
    if source_mode == "RGBA":
        Image.new("RGBA", (1200, 900), (20, 40, 80, 160)).save(source)
    else:
        source = source.with_suffix(".jpg")
        Image.new("CMYK", (1200, 900), (10, 20, 30, 0)).save(source)

    normalized = Path(ppt_tools._normalize_full_page_image(str(source)))
    with Image.open(normalized) as page:
        assert page.mode == "RGB"
        assert page.size == (1792, 1008)


def test_full_page_normalizer_applies_exif_orientation(tmp_path: Path) -> None:
    from PIL import Image

    source = tmp_path / "provider-oriented.jpg"
    image = Image.new("RGB", (640, 960), (30, 60, 90))
    exif = image.getexif()
    exif[274] = 6  # Rotate 90 degrees clockwise before crop/resize.
    image.save(source, exif=exif)

    normalized = Path(ppt_tools._normalize_full_page_image(str(source)))
    with Image.open(normalized) as page:
        assert page.mode == "RGB"
        assert page.size == (1792, 1008)


def test_full_page_assembler_rejects_missing_or_unnormalized_pages(tmp_path: Path) -> None:
    missing = ppt_tools._render_full_page_images(
        ppt_tools.parse_outline([{"title": "Missing", "image_path": str(tmp_path / "none.png")}]),
        title="Deck",
        author="DeskPet",
        output_path=str(tmp_path / "missing.pptx"),
    )
    assert missing["ok"] is False
    assert missing["error_kind"] == "missing_page_image"
    assert not (tmp_path / "missing.pptx").exists()

    raw = _provider_image(tmp_path / "raw.png")
    invalid = ppt_tools._render_full_page_images(
        ppt_tools.parse_outline([{"title": "Raw", "image_path": str(raw)}]),
        title="Deck",
        author="DeskPet",
        output_path=str(tmp_path / "invalid.pptx"),
    )
    assert invalid["ok"] is False
    assert invalid["error_kind"] == "invalid_page_image"


@pytest.mark.asyncio
async def test_full_page_preview_reuses_committed_page_images_without_com(tmp_path: Path) -> None:
    page = _page_image(tmp_path / "raw.png")
    result = ppt_tools._render_full_page_images(
        ppt_tools.parse_outline([{"title": "Page", "image_path": str(page)}]),
        title="Deck",
        author="DeskPet",
        output_path=str(tmp_path / "preview-source.pptx"),
    )
    result["output_hash"] = "deck-hash"

    previews = await nodes.render_preview(result, port=None)
    assert len(previews) == 1
    assert previews[0]["path"] == str(page.resolve())
    assert previews[0]["render_hash"] == "deck-hash"


@pytest.mark.asyncio
async def test_runtime_full_page_render_does_not_probe_or_use_legacy_renderer(
    monkeypatch, tmp_path: Path
) -> None:
    # This is a primitive-runtime unit test, not a production durable-run test.
    # Automation provides an isolated DESKPET_USER_DATA_DIR globally; without
    # clearing it PptRuntime correctly attempts to resolve a live ppt_pro run.
    monkeypatch.delenv("DESKPET_USER_DATA_DIR", raising=False)
    page = _page_image(tmp_path / "raw.png")
    output = tmp_path / "runtime.pptx"
    monkeypatch.setattr(
        ppt_tools,
        "_render_pro",
        lambda *args, **kwargs: (_ for _ in ()).throw(
            AssertionError("legacy renderer must not run")
        ),
    )
    monkeypatch.setattr(
        ppt_tools,
        "probe_image_reachable",
        lambda **kwargs: (_ for _ in ()).throw(AssertionError("probe must not run")),
    )

    result = await PptRuntime().render_ppt(
        slides=[{"title": "Page", "image_path": str(page), "notes": "note"}],
        topic="Topic",
        title="Deck",
        author="DeskPet",
        theme="minimal",
        output_path=str(output),
        render_mode="full_page_images",
        render_revision=0,
        input_hash="full-page-input",
    )
    assert result["ok"] is True
    assert output.is_file()


def test_visual_revision_groups_page_once_and_atomically_resets_image_state(monkeypatch) -> None:
    monkeypatch.setattr(ppt_tools, "_resolve_ppt_image_model", lambda: "model-a")
    records = nodes.stable_slide_records(
        [
            {"title": "One", "bullets": ["A"]},
            {"title": "Two", "bullets": ["B"]},
        ],
        full_page_images=True,
    )
    records[0]["image"].update(
        {"status": "ready", "path": "old.png", "post_hash": "old-hash", "error": "old"}
    )
    records[0]["slide"]["image_path"] = "old.png"
    untouched = nodes.content_hash(records[1])

    revised = nodes.apply_visual_revision(
        records,
        [
            {"slide_id": records[0]["slide_id"], "reason_codes": ["text_missing"]},
            {"slide_id": records[0]["slide_id"], "reason_codes": ["text_overflow"]},
        ],
    )

    image = revised[0]["image"]
    assert image["page_revision"] == 1
    assert image["revision_codes"] == ["text_missing", "text_overflow"]
    assert image["status"] == "pending"
    assert image["path"] is None and image["post_hash"] is None and image["error"] is None
    assert image["pre_hash"] != records[0]["image"]["pre_hash"]
    assert revised[0]["slide"]["image_path"] is None
    assert nodes.content_hash(revised[1]) == untouched
    assert nodes.next_pending_slide(revised)["slide_id"] == records[0]["slide_id"]


def test_layout_misfit_revision_changes_layout_and_hash(monkeypatch) -> None:
    monkeypatch.setattr(ppt_tools, "_resolve_ppt_image_model", lambda: "model-a")
    records = nodes.stable_slide_records(
        [{"title": "One", "bullets": ["A"]}, {"title": "Two", "bullets": ["B"]}],
        full_page_images=True,
    )
    before_layout = records[0]["slide"]["full_page_layout"]
    before_hash = records[0]["image"]["pre_hash"]

    revised = nodes.apply_visual_revision(
        records,
        [{"slide_id": records[0]["slide_id"], "reason_codes": ["layout_misfit"]}],
    )

    assert revised[0]["slide"]["full_page_layout"] != before_layout
    assert revised[0]["image"]["pre_hash"] != before_hash
    assert revised[0]["slide_id"] == records[0]["slide_id"]


def test_legacy_full_page_records_reuse_ready_and_replan_pending(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(ppt_tools, "_resolve_ppt_image_model", lambda: "model-a")
    records = nodes.stable_slide_records(
        [{"title": "Ready"}, {"title": "Pending"}, {"title": "Pending 2"}],
        full_page_images=True,
    )
    ready = _page_image(tmp_path / "ready.png")
    records[0]["slide"].pop("full_page_layout", None)
    records[0]["slide"]["image_path"] = str(ready)
    records[0]["image"].update(
        {"status": "ready", "path": str(ready), "post_hash": "committed"}
    )
    for record in records[1:]:
        record["slide"].pop("full_page_layout", None)
    old_pre_hash = records[0]["image"]["pre_hash"]

    normalized = nodes.normalize_full_page_records(records)

    assert normalized[0]["slide"]["full_page_layout"] == "text_left"
    assert normalized[0]["image"]["pre_hash"] == old_pre_hash
    assert normalized[0]["image"]["path"] == str(ready)
    assert all(record["slide"]["full_page_layout"] for record in normalized[1:])
    assert normalized[1]["slide"]["full_page_layout"] != "text_left"


def test_missing_ready_full_page_asset_returns_to_pending(monkeypatch, tmp_path) -> None:
    monkeypatch.setattr(ppt_tools, "_resolve_ppt_image_model", lambda: "model-a")
    records = nodes.stable_slide_records(
        [{"title": "Lost page"}], full_page_images=True
    )
    records[0]["image"].update(
        {"status": "ready", "path": str(tmp_path / "missing.png"), "post_hash": "old"}
    )
    records[0]["slide"]["image_path"] = str(tmp_path / "missing.png")

    normalized = nodes.normalize_full_page_records(records)

    assert normalized[0]["image"]["status"] == "pending"
    assert normalized[0]["image"]["path"] is None
    assert normalized[0]["slide"]["image_path"] is None


@pytest.mark.asyncio
async def test_full_page_runtime_hashes_normalized_file(monkeypatch, tmp_path: Path) -> None:
    # Keep this standalone adapter test independent from the evidence runner's
    # process-wide production userdata root.
    monkeypatch.delenv("DESKPET_USER_DATA_DIR", raising=False)
    raw = _provider_image(tmp_path / "provider.png")

    def generate(prompts, *, size, model):
        assert size == "1792x1024"
        assert model == "model-a"
        return [{"path": str(raw), "error": None}]

    monkeypatch.setattr(ppt_tools, "generate_images", generate)
    result = await PptRuntime().generate_slide_image(
        slide_id="slide-1",
        prompt="full page",
        size="1792x1024",
        input_hash="input",
        model="model-a",
        prompt_schema_version="prompt-v1",
        normalizer_version="normalizer-v1",
        page_revision=0,
        full_page=True,
    )

    final_path = Path(str(result["path"]))
    assert final_path != raw and final_path.is_file()
    assert result["output_hash"] == hashlib.sha256(final_path.read_bytes()).hexdigest()


@pytest.mark.asyncio
async def test_runtime_full_page_evaluator_passes_exact_copy(monkeypatch, tmp_path: Path) -> None:
    from deskpet.tools import ppt_visual_review

    preview = tmp_path / "preview.png"
    preview.write_bytes(b"preview")
    captured = {}

    def review(paths, metadata, *, mode):
        captured.update(paths=paths, metadata=metadata, mode=mode)
        return [
            {
                "page": 1,
                "ok": False,
                "issues": ["missing body"],
                "reason_codes": ["text_missing"],
                "action": "regenerate_page",
            }
        ]

    monkeypatch.setattr(ppt_visual_review, "review_slides", review)
    result = await PptRuntime().evaluate_ppt(
        previews=[{"path": str(preview)}],
        slides=[
            {
                "slide_id": "slide-1",
                "slide": {
                    "title": "精确标题",
                    "subtitle": "副标题",
                    "bullets": ["正文一", "正文二"],
                    "notes": "hidden note",
                },
                "image": {"mode": "full_page_images", "status": "ready"},
            }
        ],
    )

    assert captured["mode"] == "full_page_images"
    expected = captured["metadata"][0]["expected_text"]
    assert expected["title"] == "精确标题"
    assert expected["bullets"] == ["正文一", "正文二"]
    assert "notes" not in expected
    assert result["issues"][0]["slide_id"] == "slide-1"


@pytest.mark.asyncio
async def test_revision_connectivity_failure_retries_then_never_downgrades(
    monkeypatch,
) -> None:
    monkeypatch.setattr(ppt_tools, "_resolve_ppt_image_model", lambda: "model-a")
    records = nodes.stable_slide_records(
        [{"title": "One", "bullets": ["A"]}], full_page_images=True
    )
    revised = nodes.apply_visual_revision(
        records,
        [{"slide_id": records[0]["slide_id"], "reason_codes": ["text_overflow"]}],
    )

    class UnavailablePort:
        async def generate_slide_image(self, **kwargs):
            return {
                "path": None,
                "error": "provider temporarily unavailable",
                "error_kind": "connectivity",
            }

    state = {
        "values": {
            "slide_records": revised,
            "render_mode": nodes.FULL_PAGE_RENDER_MODE,
            "full_page_images": True,
            "image_fallback_reason": None,
        }
    }
    context = WorkflowContext(ports={"tool": UnavailablePort()})
    for expected_attempt in (1, 2):
        patch = await ppt_pro_graph.image_map_handler(state, context)
        values = patch.values["values"]
        assert values.get("terminal_status") is None
        assert values["slide_records"][0]["image"]["status"] == "pending"
        assert values["slide_records"][0]["image"]["provider_attempts"] == expected_attempt
        state = {"values": values}

    patch = await ppt_pro_graph.image_map_handler(state, context)
    values = patch.values["values"]

    assert values["terminal_status"] == "error"
    assert values["terminal_error"]["code"] == "ppt_full_page_generation_unavailable"
    assert values["slide_records"][0]["image"]["provider_attempts"] == 3
    assert values["render_mode"] == nodes.FULL_PAGE_RENDER_MODE
    assert "image_fallback_reason" not in values or values["image_fallback_reason"] is None


@pytest.mark.asyncio
async def test_full_page_provider_retry_recovers_same_pending_slide(
    monkeypatch, tmp_path: Path
) -> None:
    monkeypatch.setattr(ppt_tools, "_resolve_ppt_image_model", lambda: "model-a")
    records = nodes.stable_slide_records(
        [{"title": "One", "bullets": ["A"]}], full_page_images=True
    )
    ready = _page_image(tmp_path / "ready.png")

    class FlakyPort:
        calls = 0

        async def generate_slide_image(self, **kwargs):
            self.calls += 1
            if self.calls < 3:
                return {
                    "path": None,
                    "error": "transient",
                    "error_kind": "connectivity",
                }
            return {"path": str(ready)}

    state = {
        "values": {
            "slide_records": records,
            "render_mode": nodes.FULL_PAGE_RENDER_MODE,
            "full_page_images": True,
            "image_fallback_reason": None,
        }
    }
    port = FlakyPort()
    context = WorkflowContext(ports={"tool": port})
    for _ in range(3):
        patch = await ppt_pro_graph.image_map_handler(state, context)
        state = {"values": patch.values["values"]}

    image = state["values"]["slide_records"][0]["image"]
    assert port.calls == 3
    assert image["status"] == "ready"
    assert image["path"] == str(ready)
    assert state["values"]["render_mode"] == nodes.FULL_PAGE_RENDER_MODE
    assert state["values"].get("terminal_status") is None
