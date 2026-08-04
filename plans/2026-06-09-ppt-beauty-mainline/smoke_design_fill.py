"""A-4 真实模板冒烟: 高级感(01) 设计页复用,验证中文内容替换+无 lorem 残留。"""
import sys
from pathlib import Path
sys.path.insert(0, r"G:\projects\deskpet\backend")
from pptx import Presentation  # noqa: E402
from deskpet.tools.ppt_tools import ppt_create  # noqa: E402

TPL = r"G:\projects\deskpet\resources\PPT_Template\高级感 (01)---.pptx"
OUT = r"G:\projects\deskpet\.tmp\design-fill-smoke-v6.pptx"
REPORT = Path(r"G:\projects\deskpet\plans\2026-06-09-ppt-beauty-mainline\smoke_design_fill_report.txt")

outline = [
    {"layout": "title", "title": "大语言模型应用实践", "subtitle": "DeskPet 设计页复用样例"},
    {"layout": "section", "title": "一、技术背景", "subtitle": "从规则到神经网络"},
    {"layout": "bullet", "title": "核心能力", "bullets": [
        "自然语言理解与生成", "多轮对话与上下文记忆", "工具调用与代码执行", "检索增强与长期记忆"]},
    {"layout": "two_column", "title": "本地 vs 云端", "left_title": "本地部署",
     "left": ["数据不出本机", "隐私安全"], "right_title": "云端模型", "right": ["能力更强", "免维护"]},
]
res = ppt_create(outline, template=TPL, output_path=OUT, title="大语言模型应用实践")
lines = [f"ok={res.get('ok')} theme={res.get('theme')} slides={res.get('slide_count')} path={res.get('path')}"]
ok = bool(res.get("ok"))
if ok and Path(res["path"]).is_file():
    prs = Presentation(res["path"])
    lines.append(f"reopened slides={len(prs.slides)}")
    lorem_found = False
    for si, s in enumerate(prs.slides):
        texts = []
        def walk(shapes):
            global lorem_found
            for sh in shapes:
                if sh.shape_type == 6:
                    walk(sh.shapes); continue
                if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
                    t = sh.text_frame.text.replace("\n", " / ")[:60]
                    texts.append(t)
                    if "Presentations are communication" in sh.text_frame.text:
                        lorem_found = True
        walk(s.shapes)
        lines.append(f"  slide[{si}]: {texts}")
    lines.append(f"LOREM_RESIDUE={lorem_found}")
    lines.append("VERDICT=" + ("PASS" if not lorem_found and len(prs.slides) == 4 else "FAIL"))
else:
    lines.append("VERDICT=FAIL (render not ok)")
REPORT.write_text("\n".join(lines), encoding="utf-8")
print("written", REPORT)





